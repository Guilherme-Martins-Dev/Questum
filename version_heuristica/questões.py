from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from pathlib import Path
import re
import os
import shutil
import subprocess
import base64
import hashlib
import html
import mimetypes
import xml.etree.ElementTree as ET
from urllib.parse import quote
from docx.oxml.ns import qn

# =========================
# CONFIGURAÇÃO BASE
# =========================
BASE = Path(__file__).resolve().parent
SAIDA_GIFT = BASE / "banco_questoes_gift.txt"
SAIDA_XML = BASE / "banco_questoes_moodle.xml"
PASTA_IMAGENS = BASE / "imagens_questoes"

# Registro preenchido durante a leitura dos arquivos Word.
# Chave: nome único da imagem; valor: metadados e bytes.
IMAGENS_EXTRAIDAS = {}

# Configurados em escolher_formato_saida().
PREFERENCIA_SAIDA = "automatico"
FORMATO_SAIDA = ""
URL_BASE_IMAGENS = ""

TIPOS_MAPA = {
    "DTCOM": "DTCOM",
    "UNIATENEU": "Uniateneu",
}

entrada_tipo = input("Informe o tipo de material (DTCOM ou Uniateneu): ").strip().upper()
if entrada_tipo not in TIPOS_MAPA:
    raise ValueError("TIPO_MATERIAL deve ser 'DTCOM' ou 'Uniateneu'.")
TIPO_MATERIAL = TIPOS_MAPA[entrada_tipo]

DISCIPLINA = input("Informe o nome da disciplina: ").strip()
if not DISCIPLINA:
    raise ValueError("O nome da disciplina não pode ficar vazio.")


# =========================
# UTILITÁRIOS GERAIS
# =========================
def escapar_gift(texto: str) -> str:
    if not texto:
        return ""
    texto = texto.replace("\\", "\\\\")
    for ch in ["~", "=", "#", "{", "}", ":"]:
        texto = texto.replace(ch, f"\\{ch}")
    return texto.strip()


def normalizar_texto(texto: str) -> str:
    if not texto:
        return ""
    texto = texto.strip().lower()
    texto = re.sub(r"\s+", " ", texto)
    return texto


EXTENSOES_WORD_SUPORTADAS = {".docx", ".doc"}


def listar_arquivos_word(pasta: Path) -> list[Path]:
    """Lista arquivos Word suportados, ignorando temporários do Office."""
    arquivos = []
    for caminho in pasta.iterdir():
        if caminho.is_file() and caminho.suffix.lower() in EXTENSOES_WORD_SUPORTADAS and not caminho.name.startswith("~$"):
            arquivos.append(caminho)
    return sorted(arquivos, key=lambda item: item.name.lower())


def resolver_nome_arquivo_word(nome_arquivo: str, pasta: Path = BASE) -> Path | None:
    """Resolve um nome informado com ou sem a extensão .docx ou .doc."""
    nome_arquivo = nome_arquivo.strip().strip('"')
    if not nome_arquivo:
        return None
    informado = Path(nome_arquivo)
    caminho_base = informado if informado.is_absolute() else pasta / informado
    if caminho_base.suffix.lower() in EXTENSOES_WORD_SUPORTADAS:
        return caminho_base if caminho_base.exists() else None
    if caminho_base.suffix:
        return None
    for extensao in (".docx", ".doc"):
        candidato = caminho_base.with_suffix(extensao)
        if candidato.exists():
            return candidato
    return None


def _converter_doc_com_libreoffice(origem: Path, destino: Path) -> bool:
    executavel = shutil.which("soffice") or shutil.which("libreoffice")
    if not executavel:
        candidatos = [
            Path(os.environ.get("PROGRAMFILES", "")) / "LibreOffice/program/soffice.exe",
            Path(os.environ.get("PROGRAMFILES(X86)", "")) / "LibreOffice/program/soffice.exe",
        ]
        executavel = next((str(c) for c in candidatos if c.exists()), None)
    if not executavel:
        return False
    destino.parent.mkdir(parents=True, exist_ok=True)
    resultado = subprocess.run([str(executavel), "--headless", "--convert-to", "docx", "--outdir", str(destino.parent), str(origem)], capture_output=True, text=True, timeout=180, check=False)
    gerado_padrao = destino.parent / f"{origem.stem}.docx"
    if resultado.returncode == 0 and gerado_padrao.exists():
        if gerado_padrao.resolve() != destino.resolve():
            if destino.exists():
                destino.unlink()
            gerado_padrao.replace(destino)
        return True
    return False


def _converter_doc_com_word_windows(origem: Path, destino: Path) -> bool:
    """Converte via Microsoft Word usando PowerShell, sem exigir pywin32."""
    if os.name != "nt":
        return False
    powershell = shutil.which("powershell") or shutil.which("pwsh")
    if not powershell:
        return False
    destino.parent.mkdir(parents=True, exist_ok=True)
    origem_ps = str(origem.resolve()).replace("'", "''")
    destino_ps = str(destino.resolve()).replace("'", "''")
    script = (
        "$ErrorActionPreference = 'Stop'\n"
        "$word = $null\n$doc = $null\n"
        "try {\n"
        "  $word = New-Object -ComObject Word.Application\n"
        "  $word.Visible = $false\n  $word.DisplayAlerts = 0\n"
        f"  $doc = $word.Documents.Open('{origem_ps}', $false, $true)\n"
        f"  $doc.SaveAs2('{destino_ps}', 16)\n"
        "} finally {\n"
        "  if ($doc -ne $null) { $doc.Close($false) }\n"
        "  if ($word -ne $null) { $word.Quit() }\n"
        "}\n"
    )
    resultado = subprocess.run([powershell, "-NoProfile", "-NonInteractive", "-Command", script], capture_output=True, text=True, timeout=180, check=False)
    return resultado.returncode == 0 and destino.exists()


def preparar_arquivo_word(caminho: Path) -> Path:
    """Retorna um .docx pronto para leitura; converte .doc automaticamente."""
    caminho = caminho.resolve()
    extensao = caminho.suffix.lower()
    if extensao == ".docx":
        return caminho
    if extensao != ".doc":
        raise ValueError(f"Extensão não suportada: {caminho.suffix}")
    pasta_convertidos = BASE / "arquivos_doc_convertidos"
    destino = pasta_convertidos / f"{caminho.stem}.docx"
    if destino.exists() and destino.stat().st_mtime >= caminho.stat().st_mtime:
        print(f"[INFO] Reutilizando conversão existente: {destino.name}")
        return destino
    print(f"[INFO] Arquivo .doc detectado: {caminho.name}")
    print("[INFO] Convertendo automaticamente para .docx...")
    if _converter_doc_com_libreoffice(caminho, destino):
        print(f"[OK] Conversão realizada pelo LibreOffice: {destino}")
        return destino
    if _converter_doc_com_word_windows(caminho, destino):
        print(f"[OK] Conversão realizada pelo Microsoft Word: {destino}")
        return destino
    raise RuntimeError("Não foi possível converter o arquivo .doc. Instale o LibreOffice ou execute o script em um computador Windows com o Microsoft Word instalado. " f"Arquivo: {caminho}")

def romano_para_int(romano: str):
    if not romano:
        return None

    romano = romano.upper().strip()
    mapa = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}
    total = 0
    anterior = 0

    for ch in reversed(romano):
        valor = mapa.get(ch)
        if valor is None:
            return None
        if valor < anterior:
            total -= valor
        else:
            total += valor
            anterior = valor

    return total


def run_vermelho(run) -> bool:
    color = getattr(getattr(run.font, "color", None), "rgb", None)
    if color is None:
        return False
    cor = str(color).upper()
    return cor == "FF0000" or "FF0000" in cor


def run_amarelo(run) -> bool:
    highlight = getattr(run.font, "highlight_color", None)
    return highlight == WD_COLOR_INDEX.YELLOW


def run_negrito(run) -> bool:
    return bool(run.bold)


def obter_info_lista_word(paragrafo):
    """
    Retorna informações de lista automática do Word.

    Alguns arquivos .docx usam numeração automática para perguntas e
    alternativas. Nesses casos, o texto do parágrafo vem sem "1)" ou "a)",
    mas o Word registra o nível da lista:
    - ilvl == 0: item principal, normalmente a questão;
    - ilvl == 1: subitem, normalmente a alternativa.
    """
    pPr = paragrafo._p.pPr
    if pPr is None or pPr.numPr is None:
        return None, None

    numPr = pPr.numPr
    num_id = None
    ilvl = None

    if numPr.numId is not None:
        try:
            num_id = int(numPr.numId.val)
        except (TypeError, ValueError):
            num_id = None

    if numPr.ilvl is not None:
        try:
            ilvl = int(numPr.ilvl.val)
        except (TypeError, ValueError):
            ilvl = None

    return num_id, ilvl


def _extensao_por_mime(content_type: str, nome_original: str = "") -> str:
    mapa = {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/gif": ".gif",
        "image/webp": ".webp",
        "image/svg+xml": ".svg",
        "image/bmp": ".bmp",
        "image/tiff": ".tif",
        "image/x-emf": ".emf",
        "image/x-wmf": ".wmf",
    }
    if content_type in mapa:
        return mapa[content_type]
    sufixo = Path(nome_original).suffix.lower()
    if sufixo:
        return sufixo
    return mimetypes.guess_extension(content_type or "") or ".bin"


def extrair_imagens_do_paragrafo(paragrafo, doc, origem: str, indice_paragrafo: int):
    """Extrai imagens DrawingML/VML ancoradas no parágrafo e cria marcadores."""
    imagens = []
    rel_ids = []

    # Imagens modernas do Word: <a:blip r:embed="rIdX">
    for blip in paragrafo._p.xpath('.//*[local-name()="blip"]'):
        rid = blip.get(qn("r:embed"))
        if rid:
            rel_ids.append(rid)

    # Imagens legadas/VML: <v:imagedata r:id="rIdX">
    for image_data in paragrafo._p.xpath('.//*[local-name()="imagedata"]'):
        rid = image_data.get(qn("r:id"))
        if rid:
            rel_ids.append(rid)

    vistos = set()
    for ordem, rid in enumerate(rel_ids, start=1):
        if rid in vistos:
            continue
        vistos.add(rid)

        parte = doc.part.related_parts.get(rid)
        if parte is None or not hasattr(parte, "blob"):
            continue

        dados = parte.blob
        content_type = getattr(parte, "content_type", "application/octet-stream")
        nome_original = Path(str(getattr(parte, "partname", "imagem"))).name
        extensao = _extensao_por_mime(content_type, nome_original)
        digest = hashlib.sha256(dados).hexdigest()[:16]
        nome = f"img_{digest}{extensao}"
        marcador = f"__MOODLE_IMAGE_{digest.upper()}__"

        IMAGENS_EXTRAIDAS[nome] = {
            "nome": nome,
            "marcador": marcador,
            "content_type": content_type,
            "dados": dados,
            "base64": base64.b64encode(dados).decode("ascii"),
            "origem": origem,
            "indice_paragrafo": indice_paragrafo,
            "ordem": ordem,
        }
        imagens.append(IMAGENS_EXTRAIDAS[nome])

    return imagens


def extrair_paragrafos_com_runs(docx_path):
    doc = Document(docx_path)
    resultado = []

    for indice_paragrafo, p in enumerate(doc.paragraphs):
        texto = p.text.strip()
        imagens = extrair_imagens_do_paragrafo(
            p, doc, Path(docx_path).name, indice_paragrafo
        )

        if not texto and not imagens:
            continue

        runs = []
        for r in p.runs:
            runs.append(
                {
                    "text": r.text,
                    "red": run_vermelho(r),
                    "yellow": run_amarelo(r),
                    "bold": run_negrito(r),
                }
            )

        num_id, ilvl = obter_info_lista_word(p)

        # Primeiro preserva o texto do parágrafo.
        linhas = [linha.strip() for linha in texto.splitlines() if linha.strip()]
        for linha in linhas:
            resultado.append(
                {
                    "text": linha,
                    "runs": runs,
                    "alignment": p.alignment,
                    "num_id": num_id,
                    "ilvl": ilvl,
                    "images": [],
                    "paragraph_index": indice_paragrafo,
                }
            )

        # Em seguida inclui um marcador por imagem na posição correspondente.
        for imagem in imagens:
            resultado.append(
                {
                    "text": imagem["marcador"],
                    "runs": [],
                    "alignment": p.alignment,
                    "num_id": num_id,
                    "ilvl": ilvl,
                    "images": [imagem],
                    "paragraph_index": indice_paragrafo,
                }
            )

    return resultado


def detectar_alternativa_correta(paragrafo) -> bool:
    return any(
        (run["red"] or run["yellow"] or run["bold"]) and run["text"].strip()
        for run in paragrafo["runs"]
    )


def extrair_letra_correta_da_justificativa(texto: str):
    if not texto:
        return None

    padroes = [
        r"alternativa\s+[\"\'\u201c\u201d\u2018\u2019]?\s*([A-E])\s*[\"\'\u201c\u201d\u2018\u2019]?\s+est[a\u00e1]\s+correta",
        r"alternativa\s+[\"\'\u201c\u201d\u2018\u2019]?\s*([A-E])\s*[\"\'\u201c\u201d\u2018\u2019]?",
    ]

    for padrao in padroes:
        m = re.search(padrao, texto, flags=re.IGNORECASE)
        if m:
            return m.group(1).upper()

    return None


# =========================
# DETECTORES DE ESTRUTURA
# =========================
def eh_cabecalho_questao(texto: str) -> bool:
    txt = texto.strip()
    return bool(re.match(r"^Quest[aã]o\s+\d+\s*[\.\:\-]?\s*$", txt, flags=re.IGNORECASE))


def eh_cabecalho_unidade(texto: str) -> bool:
    txt = texto.strip()

    padroes = [
        r"^Unidade\s+0*\d+\b(?:\s*[:\-–—].*)?$",
        r"^Unidade\s+[IVXLCDM]+\b(?:\s*[:\-–—].*)?$",
        r"^Unidade\s+0*\d+\s*[-–]\s*Question[áa]rio\s*$",
        r"^Unidade\s+[IVXLCDM]+\s*[-–]\s*Question[áa]rio\s*$",
        r"^Bloco\s+0*\d+\b.*$",
        r"^Bloco\s+[IVXLCDM]+\b.*$",
    ]

    return any(re.match(p, txt, flags=re.IGNORECASE) for p in padroes)


def normalizar_unidade(texto: str) -> str:
    txt = texto.strip()

    m_bloco_arabico = re.match(r"^Bloco\s+0*(\d+)\b", txt, flags=re.IGNORECASE)
    if m_bloco_arabico:
        return f"Unidade {int(m_bloco_arabico.group(1))}"

    m_bloco_romano = re.match(r"^Bloco\s+([IVXLCDM]+)\b", txt, flags=re.IGNORECASE)
    if m_bloco_romano:
        valor = romano_para_int(m_bloco_romano.group(1))
        if valor is not None:
            return f"Unidade {valor}"

    m_arabico = re.search(r"(\d+)", txt)
    if m_arabico:
        return f"Unidade {int(m_arabico.group(1))}"

    m_romano = re.search(r"\b([IVXLCDM]+)\b", txt, flags=re.IGNORECASE)
    if m_romano:
        valor = romano_para_int(m_romano.group(1))
        if valor is not None:
            return f"Unidade {valor}"

    return txt


def eh_assunto(texto: str) -> bool:
    return texto.strip().lower().startswith("assunto:")


def limpar_assunto(texto: str) -> str:
    return re.sub(r"^Assunto\s*:\s*", "", texto.strip(), flags=re.IGNORECASE).strip()


def eh_titulo_forum(texto: str) -> bool:
    txt = normalizar_texto(texto)
    return (
        txt == "fórum"
        or txt == "forum"
        or bool(re.fullmatch(r"f[oó]rum\s+[ivxlcdm0-9]+", txt, flags=re.IGNORECASE))
        or (("sugest" in txt) and ("fórum" in txt or "forum" in txt))
    )


def eh_titulo_questoes(texto: str) -> bool:
    txt = normalizar_texto(texto)
    return (
        txt in {"questões", "questoes", "questionário", "questionario"}
        or txt.startswith("questões")
        or txt.startswith("questoes")
        or txt.startswith("questionário")
        or txt.startswith("questionario")
    )


def eh_titulo_gabarito(texto: str) -> bool:
    return bool(re.fullmatch(r"gabarito\s*: ?|gabarito", texto.strip(), flags=re.IGNORECASE))


def eh_linha_gabarito(texto: str) -> bool:
    txt = texto.strip()
    return bool(
        re.match(r"^(?:✅\s*)?\[?\s*Gabarito\s*\]?\s*:\s*([A-E])\s*[\)\.]?\b", txt, flags=re.IGNORECASE)
    )


def eh_linha_comentario(texto: str) -> bool:
    txt = texto.strip()
    return bool(
        re.match(
            r"^(?:💡\s*)?(?:\[?\s*(?:Feedback\s*/\s*Coment[aá]rio|Coment[aá]rio|Feedback|Justificativa)\s*\]?\s*:)",
            txt,
            flags=re.IGNORECASE,
        )
    )


def extrair_letra_gabarito_inline(texto: str):
    if not texto:
        return None

    m = re.search(r"^(?:✅\s*)?\[?\s*Gabarito\s*\]?\s*:\s*([A-E])\s*[\)\.]?\b", texto.strip(), flags=re.IGNORECASE)
    if m:
        return m.group(1).upper()

    return None


# =========================
# PADRÃO: AVALIAÇÃO + "✅ RESPOSTA CORRETA: X"
# =========================
def eh_linha_resposta_correta(texto: str) -> bool:
    """
    Detecta linhas de gabarito inline usadas em avaliações simples, por exemplo:
    - ✅ Resposta correta: B
    - Resposta correta: D
    - ✔ Resposta Correta - C

    O detector é propositalmente restrito a uma única letra entre A e E para
    evitar falsos positivos em enunciados que apenas mencionem "resposta correta".
    """
    if not texto:
        return False

    txt = texto.strip()
    return bool(
        re.match(
            r"^[✅✔☑]?\s*Resposta\s+correta\s*[:\-–—]\s*([A-E])\s*[\)\.]?\s*$",
            txt,
            flags=re.IGNORECASE,
        )
    )


def extrair_letra_resposta_correta(texto: str):
    """Extrai A-E de linhas como '✅ Resposta correta: B'."""
    if not texto:
        return None

    m = re.match(
        r"^[✅✔☑]?\s*Resposta\s+correta\s*[:\-–—]\s*([A-E])\s*[\)\.]?\s*$",
        texto.strip(),
        flags=re.IGNORECASE,
    )
    return m.group(1).upper() if m else None


def extrair_numero_questao_avaliacao(texto: str):
    """
    Extrai o número de questões escritas como:
    '1. Enunciado', '2) Enunciado' ou '18 - Enunciado'.
    """
    if not texto:
        return None

    m = re.match(r"^\s*(\d{1,3})\s*[\.\)\-–—:]\s+.+", texto.strip())
    return int(m.group(1)) if m else None


def eh_inicio_questao_avaliacao(texto: str) -> bool:
    return extrair_numero_questao_avaliacao(texto) is not None


def limpar_numero_questao_avaliacao(texto: str) -> str:
    return re.sub(
        r"^\s*\d{1,3}\s*[\.\)\-–—:]\s*",
        "",
        texto.strip(),
        count=1,
    ).strip()


def arquivo_tem_padrao_avaliacao_resposta_correta(paragrafos) -> bool:
    """
    Detecta o padrão do documento 'Avaliação – ...':
    - questão numerada no próprio texto: '1. ...';
    - alternativas explícitas: 'A) ...', 'B) ...';
    - gabarito após cada questão: '✅ Resposta correta: B'.

    Exige correspondência entre questões e linhas de resposta para não capturar
    acidentalmente outros templates suportados pelo script.
    """
    qtd_questoes = sum(
        1 for p in paragrafos if eh_inicio_questao_avaliacao(p["text"].strip())
    )
    qtd_respostas = sum(
        1 for p in paragrafos if eh_linha_resposta_correta(p["text"].strip())
    )
    qtd_alternativas = sum(
        1 for p in paragrafos if eh_alternativa_explica(p["text"].strip())
    )

    return (
        qtd_questoes >= 1
        and qtd_respostas >= 1
        and qtd_respostas == qtd_questoes
        and qtd_alternativas >= qtd_questoes * 2
    )


def separar_questoes_avaliacao_resposta_correta(paragrafos):
    """
    Separa blocos usando a numeração da questão como início e a linha
    'Resposta correta' como fechamento lógico.

    Retorna uma lista de dicionários com o número original da questão para que
    saltos de numeração (ex.: 16 -> 18) sejam preservados e auditáveis.
    """
    questoes = []
    atual = None

    for p in paragrafos:
        txt = p["text"].strip()
        if not txt:
            continue

        numero = extrair_numero_questao_avaliacao(txt)
        if numero is not None:
            if atual is not None:
                questoes.append(atual)
            atual = {
                "numero_original": numero,
                "paragrafos": [p],
            }
            continue

        if atual is None:
            # Ignora título/cabeçalho antes da primeira questão.
            continue

        atual["paragrafos"].append(p)

        if eh_linha_resposta_correta(txt):
            questoes.append(atual)
            atual = None

    if atual is not None:
        questoes.append(atual)

    return questoes


def parse_questao_avaliacao_resposta_correta(item_questao, unidade: str):
    """Converte um bloco do padrão 'Resposta correta' para a estrutura interna."""
    numero = item_questao["numero_original"]
    bloco = item_questao["paragrafos"]

    alternativas = []
    enunciado_partes = []
    letra_gabarito = None

    for idx, p in enumerate(bloco):
        txt = p["text"].strip()
        if not txt:
            continue

        if idx == 0 and eh_inicio_questao_avaliacao(txt):
            enunciado_partes.append(limpar_numero_questao_avaliacao(txt))
            continue

        if eh_linha_resposta_correta(txt):
            letra_gabarito = extrair_letra_resposta_correta(txt)
            continue

        if eh_alternativa_explica(txt):
            letra_m = re.match(r"^\s*([A-Ea-e])\s*[\)\.\-:]\s*(.+)$", txt)
            if letra_m:
                alternativas.append(
                    {
                        "letra": letra_m.group(1).upper(),
                        "texto": letra_m.group(2).strip(),
                        "correta": False,
                    }
                )
            continue

        # Permite enunciados quebrados em mais de um parágrafo.
        if not alternativas:
            enunciado_partes.append(txt)

    if letra_gabarito:
        for alt in alternativas:
            alt["correta"] = alt["letra"] == letra_gabarito

    corretas = [a for a in alternativas if a["correta"]]
    erros = []

    if not letra_gabarito:
        erros.append("linha 'Resposta correta' ausente ou inválida")
    elif not corretas:
        erros.append(
            f"gabarito {letra_gabarito} não corresponde às alternativas disponíveis "
            f"({', '.join(a['letra'] for a in alternativas) or 'nenhuma'})"
        )
    elif len(corretas) > 1:
        erros.append(f"mais de uma alternativa identificada para o gabarito {letra_gabarito}")

    enunciado_txt = "\n".join(
        parte for parte in enunciado_partes if parte and parte.strip()
    ).strip()

    incorretas = [a["texto"] for a in alternativas if not a["correta"]]

    return {
        "titulo": inferir_titulo(enunciado_txt, numero),
        "enunciado": enunciado_txt,
        "correta": corretas[0]["texto"] if len(corretas) == 1 else "",
        "incorretas": incorretas,
        "justificativa": "",
        "modo": "avaliacao_resposta_correta_inline",
        "qtd_alternativas": len(alternativas),
        "tipo": "Objetiva",
        "numero_original": numero,
        "letra_gabarito": letra_gabarito,
        "erros": erros,
    }


def processar_avaliacao_resposta_correta(caminho_arquivo, unidade_padrao=None):
    """
    Processa integralmente avaliações cujo gabarito vem após cada questão.

    Diferentemente dos parsers legados, este modo NUNCA assume a primeira
    alternativa como correta. Questões inconsistentes são ignoradas com aviso,
    impedindo a geração silenciosa de um gabarito incorreto.
    """
    paragrafos = extrair_paragrafos_com_runs(caminho_arquivo)
    unidade = extrair_unidade_do_nome_arquivo(caminho_arquivo) or unidade_padrao or "Unidade 1"
    itens = separar_questoes_avaliacao_resposta_correta(paragrafos)
    saida_blocos = []

    numeros = [item["numero_original"] for item in itens]
    if numeros:
        esperados = set(range(min(numeros), max(numeros) + 1))
        ausentes = sorted(esperados - set(numeros))
        if ausentes:
            print(
                f"[AVISO] {unidade}: numeração de questões com lacunas. "
                f"Ausentes: {', '.join(map(str, ausentes))}."
            )

    for item in itens:
        numero = item["numero_original"]
        q = parse_questao_avaliacao_resposta_correta(item, unidade)

        if q["erros"]:
            print(
                f"[AVISO] {unidade} | Questão {numero:02d}: "
                + "; ".join(q["erros"])
            )

        if q["enunciado"] and q["correta"] and 2 <= q["qtd_alternativas"] <= 5:
            saida_blocos.append(montar_gift(unidade, q))
            print(
                f"[OK] {unidade} | Questão {numero:02d} | "
                f"modo={q['modo']} | alternativas={q['qtd_alternativas']} | "
                f"gabarito={q['letra_gabarito']}"
            )
        else:
            print(
                f"[AVISO] {unidade} | Questão {numero:02d}: ignorada | "
                f"alternativas={q['qtd_alternativas']} | "
                f"gabarito={q['letra_gabarito'] or 'não identificado'}"
            )

    return saida_blocos


def limpar_linha_comentario(texto: str) -> str:
    txt = texto.strip()
    txt = re.sub(
        r"^(?:💡\s*)?\[?\s*(?:Feedback\s*/\s*Coment[aá]rio|Coment[aá]rio|Feedback|Justificativa)\s*\]?\s*:\s*",
        "",
        txt,
        flags=re.IGNORECASE,
    ).strip()
    txt = re.sub(r"^💡\s*", "", txt).strip()
    return txt


def extrair_gabaritos_de_tabelas(docx_path):
    doc = Document(docx_path)
    gabaritos = []

    for table in doc.tables:
        textos = []
        for row in table.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if txt:
                    textos.append(txt)

        if not any(normalizar_texto(t) == "gabarito" for t in textos):
            continue

        numeros = []
        letras = []

        for txt in textos:
            txt_limpo = txt.strip()
            if re.fullmatch(r"\d{1,2}", txt_limpo):
                numeros.append(int(txt_limpo))
            elif re.fullmatch(r"[A-Ea-e]", txt_limpo):
                letras.append(txt_limpo.upper())

        if numeros and letras:
            gabaritos.append({num: letra for num, letra in zip(numeros, letras)})

    return gabaritos


def eh_inicio_item_numerado(texto: str) -> bool:
    return bool(re.match(r"^\d+\s*[-–\.\)]\s+.+", texto.strip()))


def eh_inicio_questao_numerada(texto: str) -> bool:
    return eh_inicio_item_numerado(texto)


def limpar_numero_item_inicial(texto: str) -> str:
    return re.sub(r"^\d+\s*[-–\.\)]\s*", "", texto.strip())


def limpar_numero_questao_inicial(texto: str) -> str:
    return limpar_numero_item_inicial(texto)


def eh_cabecalho_bloco_interno(texto: str) -> bool:
    txt = texto.strip()
    return bool(re.match(r"^Bloco\s+(?:0*\d+|[IVXLCDM]+)\b.*$", txt, flags=re.IGNORECASE))


def limpar_titulo_bloco_interno(texto: str) -> str:
    txt = texto.strip()
    txt = re.sub(r"^Bloco\s+(?:0*\d+|[IVXLCDM]+)\s*[-–:]?\s*", "", txt, flags=re.IGNORECASE).strip()
    txt = re.sub(r"\s*\(Quest(?:ões|oes|ão|ao).*?\)\s*$", "", txt, flags=re.IGNORECASE).strip()
    return txt


def extrair_unidade_do_nome_arquivo(caminho_arquivo) -> str | None:
    """
    Extrai a unidade a partir do nome do arquivo.

    Correção aplicada na auditoria:
    os arquivos enviados usam o padrão abreviado "UNI 01", "UNI 02" etc.,
    enquanto a versão anterior reconhecia apenas "UNIDADE 01".
    Agora os dois padrões são aceitos:
    - UNIDADE 01 / UNIDADE I
    - UNI 01 / UNI I
    - UNID 01 / UNID I
    """
    nome = Path(caminho_arquivo).name

    m = re.search(r"\bUNI(?:DADE|D\.?)?\s*0*(\d+)\b", nome, flags=re.IGNORECASE)
    if m:
        return f"Unidade {int(m.group(1))}"

    m = re.search(r"\bUNI(?:DADE|D\.?)?\s*([IVXLCDM]+)\b", nome, flags=re.IGNORECASE)
    if m:
        valor = romano_para_int(m.group(1))
        if valor is not None:
            return f"Unidade {valor}"

    return None

def eh_subtitulo_tematica_numerado(texto: str) -> bool:
    """
    Detecta linhas como "1. Conceito de psicomotricidade" que funcionam
    como subtítulo/tema e não como questão.

    Regra prática:
    - começa com numeração simples (1., 2., 3. ...);
    - não termina com dois-pontos nem interrogação;
    - não contém comando/enunciado típico de questão;
    - é curta o bastante para parecer título temático.
    """
    txt = texto.strip()
    m = re.match(r"^(\d{1,2})\s*[\.\-–)]\s+(.+)$", txt)
    if not m:
        return False

    resto = m.group(2).strip()
    resto_norm = normalizar_texto(resto)

    if not resto or len(resto.split()) > 12:
        return False

    if resto.endswith(":") or resto.endswith("?"):
        return False

    gatilhos_enunciado = [
        "assinale", "marque", "identifique", "indique", "avalie",
        "considere", "qual", "quais", "pode ser definida", "pode ser definido",
        "refere-se", "corresponde", "implica", "está relacionada",
    ]
    if any(g in resto_norm for g in gatilhos_enunciado):
        return False

    return True


def eh_justificativa(texto: str) -> bool:
    return eh_linha_comentario(texto)


def eh_subitem_interno(texto: str) -> bool:
    """
    Detecta linhas que são claramente subitens internos de uma questão:
    - Algarismos romanos curtos: I –, II –, III –, IV –, V –
    - Letras maiúsculas: A –, B –, C –, D –
    - Parênteses: ( ) –, (  ) –
    Essas linhas NÃO devem ser tratadas como início de nova questão.

    Observação importante:
    siglas como "MDIC -" ou "CIF -" não devem ser confundidas com numerais
    romanos. Por isso, o trecho romano só é aceito quando representa um valor
    pequeno e plausível para enumeração interna.
    """
    txt = texto.strip()

    m_romano = re.match(r"^([IVXLCDM]{1,5})\s*[-–\.]\s+.+", txt, flags=re.IGNORECASE)
    if m_romano:
        valor = romano_para_int(m_romano.group(1))
        if valor is not None and 1 <= valor <= 20:
            return True

    # A -, B -, C -, D - (letra maiúscula isolada + separador, sem ser alternativa a)/b))
    if re.match(r"^[A-D]\s*[-–]\s+.+", txt):
        return True

    # ( ) texto  ou  (  ) texto  — assertivas de V/F
    if re.match(r"^\(\s*\)\s*[-–]?\s+.+", txt):
        return True

    return False

def eh_sequencia_vf(texto: str) -> bool:
    txt = re.sub(r"\s+", " ", texto.strip().upper())
    return bool(
        re.fullmatch(r"[VF](?:\s*,\s*[VF])+(?:\s*E\s*[VF])?", txt)
        or re.fullmatch(r"[VF](?:\s*E\s*[VF])+", txt)
    )


def eh_linha_comando(texto: str) -> bool:
    txt = normalizar_texto(texto)

    gatilhos = [
        "assinale a alternativa",
        "agora, assinale a alternativa",
        "agora assinale a alternativa",
        "assinale entre as alternativas",
        "assinale, entre as alternativas",
        "assinale a alternativa abaixo",
        "assinale a alternativa que apresenta",
        "agora, assinale a alternativa que apresenta a sequência correta",
        "agora assinale a alternativa que apresenta a sequência correta",
        "agora, assinale a alternativa que representa a sequência correta",
        "agora assinale a alternativa que representa a sequência correta",
        "está correto apenas o que se afirma em",
        "é correto afirmar",
        "podemos afirmar que",
        "marque a alternativa",
        "marque a opção",
        "identifique a alternativa",
        "identifique entre as opções",
        "identifique entre as alternativas",
        "identifique, entre as alternativas",
        "indique a alternativa",
        "indique entre as alternativas",
        "avalie as alternativas",
        "avalie as afirmativas",
        "considerando o texto acima",
        "diante desse cenário",
        "de acordo com o trecho",
        "de acordo com os conteúdos estudados",
        "a seguir",
        "qual das alternativas",
        "escolha a alternativa correta",
        "marque (v) para verdadeiro",
    ]

    if any(g in txt for g in gatilhos):
        return True

    if (txt.endswith(":") or txt.endswith("?")) and any(
        verbo in txt
        for verbo in [
            "assinale",
            "marque",
            "identifique",
            "indique",
            "avalie",
            "considere",
            "podemos afirmar",
            "é correto afirmar",
            "qual",
            "escolha",
        ]
    ):
        return True

    return False


def eh_alternativa_explica(texto: str) -> bool:
    return bool(re.match(r"^\s*[A-Ea-e][\)\.\-:]\s*.+", texto))


def limpar_marcador_alternativa(texto: str) -> str:
    return re.sub(r"^\s*[A-Ea-e][\)\.\-:]\s*", "", texto).strip()


def eh_ruido_formulario(texto: str) -> bool:
    txt = normalizar_texto(texto)
    return txt in {
        "parte superior do formulário",
        "parte superior do formulario",
        "parte inferior do formulário",
        "parte inferior do formulario",
    }


def normalizar_texto_identificacao_bloco(bloco, idx: int) -> str:
    txt_original = bloco[idx]["text"].strip()
    txt = txt_original

    if eh_ruido_formulario(txt):
        return ""

    if idx > 0:
        anterior = bloco[idx - 1]["text"].strip()
        if re.match(r"^([A-Ea-e])\s*Parte\s+(superior|inferior)\s+do\s+formul[áa]rio\s*$", anterior, flags=re.IGNORECASE):
            if txt.startswith((")", ".", "-", ":")):
                return ""

    txt = re.sub(r"\s*Parte\s+(superior|inferior)\s+do\s+formul[áa]rio\s*", "", txt, flags=re.IGNORECASE).strip()

    if re.match(r"^[A-Ea-e]$", txt) and idx + 1 < len(bloco):
        prox = bloco[idx + 1]["text"].strip()
        if prox.startswith((")", ".", "-", ":")):
            txt = txt + prox

    m_quebrado = re.match(r"^([A-Ea-e])\s*Parte\s+(superior|inferior)\s+do\s+formul[áa]rio\s*$", txt_original, flags=re.IGNORECASE)
    if m_quebrado and idx + 1 < len(bloco):
        prox = bloco[idx + 1]["text"].strip()
        if prox.startswith((")", ".", "-", ":")):
            txt = m_quebrado.group(1) + prox

    return txt.strip()


def inferir_titulo(enunciado: str, numero: int, assunto: str = "") -> str:
    if assunto:
        assunto = re.sub(r"\s+", " ", assunto).strip()
        if len(assunto) > 60:
            assunto = assunto[:60].rsplit(" ", 1)[0]
        return f"Questão {numero:02d} - {assunto}"

    linhas = [l.strip() for l in enunciado.splitlines() if l.strip()]
    primeira = linhas[0] if linhas else f"Questão {numero}"
    primeira = re.sub(r"\s+", " ", primeira)

    if len(primeira) > 60:
        primeira = primeira[:60].rsplit(" ", 1)[0]

    return f"Questão {numero:02d} - {primeira}"


# =========================
# MODO PADRÃO ANTIGO
# =========================
def separar_questoes(paragrafos):
    questoes = []
    bloco = []

    for p in paragrafos:
        txt = p["text"].strip()

        if eh_cabecalho_questao(txt):
            if bloco:
                questoes.append(bloco)
            bloco = [p]
        else:
            if bloco:
                bloco.append(p)

    if bloco:
        questoes.append(bloco)

    return questoes


def indice_justificativa(bloco):
    for i, p in enumerate(bloco):
        if eh_justificativa(p["text"].strip()):
            return i
    return None


def localizar_por_comando(bloco):
    """
    Localiza o comando que antecede o bloco de alternativas.

    A busca é interrompida ao encontrar a primeira alternativa explícita.
    Isso impede que expressões existentes dentro das alternativas, como
    "a seguir", sejam confundidas com um novo comando da questão.
    """
    ultimo_comando_idx = None

    for i, p in enumerate(bloco):
        txt = p["text"].strip()

        # Depois que as alternativas começam, nenhuma linha deve substituir
        # o comando já localizado. Esse era o motivo de questões com textos
        # como "tendo sempre um padrão a seguir" perderem as opções a) e b).
        if eh_alternativa_explica(txt):
            break

        if eh_justificativa(txt) or eh_cabecalho_questao(txt):
            if i > 0:
                break
            continue

        if eh_linha_comando(txt):
            ultimo_comando_idx = i

    if ultimo_comando_idx is None:
        return None

    inicio = ultimo_comando_idx + 1
    fim = len(bloco)

    for i in range(inicio, len(bloco)):
        txt = bloco[i]["text"].strip()
        if eh_justificativa(txt) or eh_cabecalho_questao(txt):
            fim = i
            break

    if inicio < fim:
        return inicio, fim, "por_comando"

    return None


def localizar_bloco_explicito_antes_da_justificativa(bloco):
    idx_just = indice_justificativa(bloco)
    if idx_just is None:
        return None

    alt_indices = []
    for i in range(idx_just):
        txt = bloco[i]["text"].strip()
        if eh_alternativa_explica(txt):
            alt_indices.append(i)

    if len(alt_indices) < 4:
        return None

    grupos = []
    grupo = [alt_indices[0]]

    for i in range(1, len(alt_indices)):
        if alt_indices[i] == alt_indices[i - 1] + 1:
            grupo.append(alt_indices[i])
        else:
            grupos.append(grupo)
            grupo = [alt_indices[i]]

    grupos.append(grupo)
    grupos_validos = [g for g in grupos if len(g) >= 4]

    if not grupos_validos:
        return None

    grupo_final = grupos_validos[-1]
    return grupo_final[0], grupo_final[-1] + 1, "por_bloco_explicito"


def localizar_por_cauda(bloco):
    idx_just = indice_justificativa(bloco)
    if idx_just is None:
        return None

    candidatos = []
    for i in range(idx_just - 1, -1, -1):
        txt = bloco[i]["text"].strip()
        if not txt:
            continue
        if eh_cabecalho_questao(txt):
            break
        if eh_justificativa(txt):
            continue
        candidatos.append(i)
        if len(candidatos) == 5:
            break

    candidatos = sorted(candidatos)

    if len(candidatos) == 5:
        return candidatos[0], candidatos[-1] + 1, "por_cauda"

    return None


def localizar_regiao_de_alternativas(bloco):
    for estrategia in (
        localizar_por_comando,
        localizar_bloco_explicito_antes_da_justificativa,
        localizar_por_cauda,
    ):
        regiao = estrategia(bloco)
        if regiao is not None:
            return regiao

    return None


def coletar_alternativas_modo_explicito(bloco, inicio, fim):
    alternativas = []

    for i in range(inicio, fim):
        txt = bloco[i]["text"].strip()
        if eh_alternativa_explica(txt):
            alternativas.append(
                {
                    "texto": limpar_marcador_alternativa(txt),
                    "correta": detectar_alternativa_correta(bloco[i]),
                    "indice": i,
                }
            )

    return alternativas


def coletar_alternativas_modo_posicional(bloco, inicio, fim):
    alternativas = []

    for i in range(inicio, fim):
        txt = bloco[i]["text"].strip()

        if not txt:
            continue

        if eh_justificativa(txt) or eh_cabecalho_questao(txt):
            break

        if eh_linha_comando(txt):
            continue

        alternativas.append(
            {
                "texto": limpar_marcador_alternativa(txt),
                "correta": detectar_alternativa_correta(bloco[i]),
                "indice": i,
            }
        )

    if len(alternativas) > 5:
        alternativas = alternativas[:5]

    return alternativas


def extrair_justificativa(bloco):
    partes = []
    capturando = False

    for p in bloco:
        txt = p["text"].strip()

        if eh_justificativa(txt):
            capturando = True
            txt = re.sub(
                r"^(Justificativa|Feedback|Coment[aá]rio)\s*:?\s*",
                "",
                txt,
                flags=re.IGNORECASE,
            )
            partes.append(txt)
            continue

        if capturando:
            if eh_cabecalho_questao(txt):
                break
            partes.append(txt)

    return " ".join(partes).strip()


def parse_questao_modo_antigo(bloco, numero, unidade):
    justificativa = extrair_justificativa(bloco)
    regiao_info = localizar_regiao_de_alternativas(bloco)

    alternativas = []
    modo = "nao_identificado"
    alt_indices = set()

    if regiao_info is not None:
        inicio, fim, estrategia = regiao_info

        alternativas = coletar_alternativas_modo_explicito(bloco, inicio, fim)

        if len(alternativas) < 4:
            alternativas = coletar_alternativas_modo_posicional(bloco, inicio, fim)
            modo = f"fallback_posicional_{estrategia}"
        else:
            modo = f"explicito_regex_{estrategia}"

        alt_indices = {a["indice"] for a in alternativas}

    enunciado_partes = []
    for i, p in enumerate(bloco):
        txt = p["text"].strip()

        if i == 0 and eh_cabecalho_questao(txt):
            continue

        if eh_justificativa(txt):
            break

        if i in alt_indices:
            continue

        enunciado_partes.append(txt)

    enunciado_txt = "\n".join([x for x in enunciado_partes if x.strip()]).strip()

    corretas = [a for a in alternativas if a["correta"]]

    if not corretas and alternativas:
        letra = extrair_letra_correta_da_justificativa(justificativa)
        if letra:
            idx = ord(letra) - ord("A")
            if 0 <= idx < len(alternativas):
                alternativas[idx]["correta"] = True
                corretas = [alternativas[idx]]

    if not corretas and alternativas:
        alternativas[0]["correta"] = True
        corretas = [alternativas[0]]
        print(
            f"[AVISO] {unidade} | Questão {numero:02d}: "
            f"correta não encontrada por cor/justificativa. Assumindo primeira alternativa."
        )

    incorretas = [a for a in alternativas if not a["correta"]]
    titulo = inferir_titulo(enunciado_txt, numero)

    return {
        "titulo": titulo,
        "enunciado": enunciado_txt,
        "correta": corretas[0]["texto"] if corretas else "",
        "incorretas": [a["texto"] for a in incorretas],
        "justificativa": justificativa,
        "modo": modo,
        "qtd_alternativas": len(alternativas),
        "tipo": "Objetiva",
    }


# =========================
# MODO ARQUIVO ÚNICO
# =========================
def eh_linha_alternativa_implicita(texto: str) -> bool:
    txt = texto.strip()

    if not txt:
        return False

    if eh_cabecalho_unidade(txt) or eh_assunto(txt) or eh_justificativa(txt) or eh_linha_gabarito(txt):
        return False

    if eh_titulo_forum(txt) or eh_titulo_questoes(txt):
        return False

    if eh_inicio_questao_numerada(txt) or eh_cabecalho_questao(txt):
        return False

    if eh_linha_comando(txt):
        return False

    # subitens internos (I-, II-, A-, ( )) NÃO são alternativas finais
    if eh_subitem_interno(txt):
        return False

    if txt.endswith(":") or txt.endswith("?"):
        return False

    if eh_alternativa_explica(txt):
        return True

    if eh_sequencia_vf(txt):
        return True

    palavras = txt.split()
    qtd_palavras = len(palavras)

    # Assertivas longas terminadas em ponto e vírgula normalmente fazem parte
    # do enunciado, não do conjunto final de alternativas.
    if txt.endswith(";") and qtd_palavras > 8:
        return False

    if 1 <= qtd_palavras <= 35 and len(txt) <= 320:
        return True

    return False

def contar_alternativas_consecutivas(paragrafos_unidade, start_idx):
    count = 0
    i = start_idx

    while i < len(paragrafos_unidade):
        txt = paragrafos_unidade[i]["text"].strip()

        if not txt:
            i += 1
            continue

        if eh_linha_alternativa_implicita(txt):
            count += 1
            i += 1
            continue

        break

    return count


def eh_inicio_questao_sem_numero(paragrafos_unidade, idx, bloco_aberto: bool = False) -> bool:
    """
    Detecta o início de uma questão objetiva sem numeração explícita.

    Regras:
    - No início da seção, aceita uma questão sem número quando a linha atual
      funciona como enunciado e é seguida por 4+ alternativas consecutivas.
    - Com bloco já aberto, só aceita a abertura de nova questão se as próximas
      alternativas explícitas forem consecutivas e imediatas, evitando atravessar
      a fronteira da próxima questão.
    - Linhas que já são alternativas do tipo V/F (ex.: "F, V e V") nunca devem
      iniciar uma nova questão.
    """
    txt = paragrafos_unidade[idx]["text"].strip()

    if not txt:
        return False

    if eh_cabecalho_unidade(txt) or eh_assunto(txt) or eh_justificativa(txt) or eh_linha_gabarito(txt):
        return False

    if eh_titulo_forum(txt) or eh_titulo_questoes(txt):
        return False

    if eh_inicio_questao_numerada(txt) or eh_cabecalho_questao(txt):
        return False

    if eh_alternativa_explica(txt) or eh_sequencia_vf(txt):
        return False

    if eh_linha_comando(txt):
        return False

    if eh_subitem_interno(txt):
        return False

    prox_idx = idx + 1
    qtd_opcoes = contar_alternativas_consecutivas(paragrafos_unidade, prox_idx)

    if bloco_aberto:
        qtd_explicitas = 0
        j = prox_idx

        while j < len(paragrafos_unidade):
            prox_txt = paragrafos_unidade[j]["text"].strip()

            if not prox_txt:
                j += 1
                continue

            if eh_alternativa_explica(prox_txt):
                qtd_explicitas += 1
                j += 1
                continue

            break

        return qtd_explicitas >= 4

    if txt.endswith("?"):
        return False

    return qtd_opcoes >= 4


def separar_unidades_documento_unico(paragrafos):
    unidades = []
    unidade_atual = None
    bloco_atual = []

    for p in paragrafos:
        txt = p["text"].strip()

        if eh_cabecalho_unidade(txt):
            if unidade_atual is not None:
                unidades.append((normalizar_unidade(unidade_atual), bloco_atual))
            unidade_atual = txt
            bloco_atual = []
        else:
            if unidade_atual is not None:
                bloco_atual.append(p)

    if unidade_atual is not None:
        unidades.append((normalizar_unidade(unidade_atual), bloco_atual))

    return unidades


def dividir_secao_unidade(bloco_unidade):
    forum = []
    questoes = []
    secao = "questoes"

    for p in bloco_unidade:
        txt = p["text"].strip()

        if eh_titulo_gabarito(txt):
            secao = "gabarito"
            continue

        if eh_titulo_forum(txt):
            secao = "forum"
            continue

        if eh_titulo_questoes(txt):
            secao = "questoes"
            continue

        if secao == "questoes":
            questoes.append(p)
        elif secao == "forum":
            forum.append(p)

    return forum, questoes


def separar_forum_documento_unico(paragrafos_forum):
    blocos = []
    bloco = []
    encontrou_numeracao = any(eh_inicio_item_numerado(p["text"].strip()) for p in paragrafos_forum)

    if encontrou_numeracao:
        for p in paragrafos_forum:
            txt = p["text"].strip()

            if not txt:
                continue

            if eh_inicio_item_numerado(txt):
                if bloco:
                    blocos.append(bloco)
                bloco = [p]
            else:
                if bloco:
                    bloco.append(p)

        if bloco:
            blocos.append(bloco)

        return blocos

    bloco_unico = [p for p in paragrafos_forum if p["text"].strip()]
    if bloco_unico:
        blocos.append(bloco_unico)

    return blocos


def parse_questao_aberta_forum(bloco, numero, unidade):
    partes = []

    for i, p in enumerate(bloco):
        txt = p["text"].strip()
        if i == 0 and eh_inicio_item_numerado(txt):
            txt = limpar_numero_item_inicial(txt)
        partes.append(txt)

    enunciado_txt = "\n".join([x for x in partes if x.strip()]).strip()
    titulo = inferir_titulo(enunciado_txt, numero)

    return {
        "titulo": titulo,
        "enunciado": enunciado_txt,
        "correta": "",
        "incorretas": [],
        "justificativa": "",
        "modo": "forum_aberto",
        "qtd_alternativas": 0,
        "tipo": "Discursiva",
    }


def linha_parece_pergunta_ou_comando(texto: str) -> bool:
    txt = texto.strip()

    if not txt:
        return False

    if eh_linha_comando(txt):
        return True

    if txt.endswith(":") or txt.endswith("?"):
        return True

    # Alguns enunciados vêm como definição/afirmação longa finalizada com ponto,
    # seguida diretamente pelas alternativas. Nesses casos, também devemos
    # aceitá-la como cabeçalho da questão.
    if (
        txt.endswith(".")
        and not eh_alternativa_explica(txt)
        and not eh_subitem_interno(txt)
        and 12 <= len(txt.split()) <= 60
    ):
        return True

    return False


def localizar_fim_bloco_alternativas_unico(bloco, start_idx):
    """
    Localiza um bloco consecutivo de 4 ou 5 alternativas.

    Ajuste restrito à etapa de identificação:
    - prioriza alternativas explícitas consecutivas (a), b), c), d));
    - ignora ruídos "Parte superior/inferior do formulário";
    - recompõe marcador quebrado em casos como "bParte superior..." + ") texto".

    Se não encontrar bloco explícito válido, preserva o comportamento antigo
    com a heurística implícita.
    """
    if start_idx <= 0 or start_idx >= len(bloco):
        return None

    txt_inicio = normalizar_texto_identificacao_bloco(bloco, start_idx)
    txt_anterior = normalizar_texto_identificacao_bloco(bloco, start_idx - 1)

    if linha_parece_pergunta_ou_comando(txt_anterior):
        indices_explicitos = []
        i = start_idx

        while i < len(bloco):
            txt_atual = normalizar_texto_identificacao_bloco(bloco, i)

            if not txt_atual:
                i += 1
                continue

            if eh_alternativa_explica(txt_atual):
                indices_explicitos.append(i)
                i += 1
                continue

            break

        if len(indices_explicitos) in {4, 5}:
            return indices_explicitos[-1]

    if not eh_linha_alternativa_implicita(txt_inicio):
        return None

    if not linha_parece_pergunta_ou_comando(txt_anterior):
        return None

    indices = [start_idx]
    i = start_idx + 1
    while i < len(bloco):
        prox_txt = normalizar_texto_identificacao_bloco(bloco, i)
        if not prox_txt:
            i += 1
            continue
        if eh_linha_alternativa_implicita(prox_txt):
            indices.append(i)
            i += 1
            continue
        break

    tamanho = len(indices)
    if 4 <= tamanho <= 5:
        return indices[-1]

    if tamanho > 5:
        idx_quinta = indices[4]
        txt_quinta = normalizar_texto_identificacao_bloco(bloco, idx_quinta)
        if linha_parece_pergunta_ou_comando(txt_quinta):
            return indices[3]

        idx_sexta = indices[5] if tamanho >= 6 else None
        if idx_sexta is not None:
            txt_sexta = normalizar_texto_identificacao_bloco(bloco, idx_sexta)
            if linha_parece_pergunta_ou_comando(txt_sexta):
                return indices[4]

    return None


def separar_questoes_documento_unico(paragrafos_unidade):
    """
    Separa questões do documento único.

    Prioridade:
    1) Se houver numeração explícita (01., 02., 03.), usa a numeração como
       fronteira principal, o que é mais estável para arquivos de questionário.
    2) Se não houver numeração, mantém a heurística baseada no fechamento do
       bloco de alternativas.

    Observação importante:
    - Em arquivos separados por unidade, cabeçalhos como "Bloco 1 – ..." são
      apenas divisões temáticas internas e NÃO novas unidades.
    """
    conteudo = []
    assunto_atual = ""

    for p in paragrafos_unidade:
        txt = p["text"].strip()

        if not txt:
            continue

        if eh_assunto(txt):
            assunto_atual = limpar_assunto(txt)
            continue

        if eh_titulo_forum(txt) or eh_titulo_questoes(txt) or eh_titulo_gabarito(txt):
            continue

        conteudo.append((assunto_atual, p))

    # Correção aplicada na auditoria:
    # arquivos no template Word tradicional usam "Questão 1", "Questão 2"
    # como fronteira principal. Sem essa etapa, o parser podia confundir
    # enumerações internas "1.", "2.", "3." com início de novas questões.
    if any(eh_cabecalho_questao(p["text"].strip()) for _, p in conteudo):
        questoes = []
        bloco = []
        assunto_bloco = ""
        assunto_pendente = ""

        for assunto, p in conteudo:
            txt = p["text"].strip()

            if eh_cabecalho_bloco_interno(txt):
                assunto_pendente = limpar_titulo_bloco_interno(txt)
                continue

            if eh_cabecalho_questao(txt):
                if bloco:
                    questoes.append((assunto_bloco, bloco))

                bloco = [p]

                if assunto:
                    assunto_bloco = assunto
                elif assunto_pendente:
                    assunto_bloco = assunto_pendente
                else:
                    assunto_bloco = ""

                assunto_pendente = ""
                continue

            if bloco:
                bloco.append(p)

        if bloco:
            questoes.append((assunto_bloco, bloco))

        return questoes

    if any(eh_inicio_questao_numerada(p["text"].strip()) and not eh_subtitulo_tematica_numerado(p["text"].strip()) for _, p in conteudo):
        questoes = []
        bloco = []
        assunto_bloco = ""
        assunto_pendente = ""

        for assunto, p in conteudo:
            txt = p["text"].strip()

            if eh_cabecalho_bloco_interno(txt):
                assunto_pendente = limpar_titulo_bloco_interno(txt)
                continue

            if eh_subtitulo_tematica_numerado(txt):
                assunto_pendente = limpar_numero_item_inicial(txt)
                continue

            if eh_inicio_questao_numerada(txt):
                if bloco:
                    questoes.append((assunto_bloco, bloco))

                bloco = [p]

                if assunto:
                    assunto_bloco = assunto
                elif assunto_pendente:
                    assunto_bloco = assunto_pendente
                else:
                    assunto_bloco = ""

                assunto_pendente = ""
            else:
                if bloco:
                    bloco.append(p)

        if bloco:
            questoes.append((assunto_bloco, bloco))

        return questoes

    questoes = []
    bloco = []
    assunto_bloco = ""
    assunto_pendente = ""

    for idx, (assunto, p) in enumerate(conteudo):
        txt = p["text"].strip()

        if eh_cabecalho_bloco_interno(txt):
            assunto_pendente = limpar_titulo_bloco_interno(txt)
            continue

        if eh_subtitulo_tematica_numerado(txt):
            assunto_pendente = limpar_numero_item_inicial(txt)
            continue

        if not bloco:
            if (
                not eh_alternativa_explica(txt)
                and not eh_linha_gabarito(txt)
                and not eh_linha_comentario(txt)
            ):
                bloco = [p]
                if assunto:
                    assunto_bloco = assunto
                elif assunto_pendente:
                    assunto_bloco = assunto_pendente
                else:
                    assunto_bloco = ""
                assunto_pendente = ""
            continue

        bloco.append(p)

        if eh_linha_comentario(txt):
            questoes.append((assunto_bloco, bloco))
            bloco = []
            assunto_bloco = ""
            continue

        if eh_linha_gabarito(txt):
            prox_txt = ""
            for _, prox in conteudo[idx + 1:]:
                prox_txt = prox["text"].strip()
                if prox_txt:
                    break
            if not prox_txt or eh_cabecalho_bloco_interno(prox_txt) or eh_subtitulo_tematica_numerado(prox_txt):
                questoes.append((assunto_bloco, bloco))
                bloco = []
                assunto_bloco = ""

    if bloco:
        questoes.append((assunto_bloco, bloco))

    return questoes

def coletar_alternativas_explicitas_unico(bloco):
    alternativas = []
    bloco_atual = []

    for i, p in enumerate(bloco):
        txt = normalizar_texto_identificacao_bloco(bloco, i)

        if not txt:
            continue

        if eh_alternativa_explica(txt):
            bloco_atual.append(
                {
                    "texto": limpar_marcador_alternativa(txt),
                    "correta": detectar_alternativa_correta(p),
                    "indice": i,
                }
            )
            continue

        if bloco_atual:
            if len(bloco_atual) in {4, 5}:
                return bloco_atual
            bloco_atual = []

    if len(bloco_atual) in {4, 5}:
        return bloco_atual

    return alternativas


def coletar_alternativas_explicitas_flexivel_unico(bloco, minimo=2):
    alternativas = []
    bloco_atual = []

    for i, p in enumerate(bloco):
        txt = normalizar_texto_identificacao_bloco(bloco, i)

        if not txt:
            continue

        if eh_alternativa_explica(txt):
            bloco_atual.append(
                {
                    "texto": limpar_marcador_alternativa(txt),
                    "correta": detectar_alternativa_correta(p),
                    "indice": i,
                }
            )
            continue

        if bloco_atual:
            if minimo <= len(bloco_atual) <= 5:
                return bloco_atual
            bloco_atual = []

    if minimo <= len(bloco_atual) <= 5:
        return bloco_atual

    return alternativas


def coletar_alternativas_por_janela_unico(bloco):
    melhor_bloco = []

    for inicio in range(1, len(bloco)):
        txt_inicio = bloco[inicio]["text"].strip()

        if not eh_linha_alternativa_implicita(txt_inicio):
            continue

        candidatos = []
        i = inicio

        while i < len(bloco):
            txt = bloco[i]["text"].strip()

            if not txt:
                i += 1
                continue

            if eh_linha_alternativa_implicita(txt):
                candidatos.append(i)
                i += 1
                continue

            break

        if 4 <= len(candidatos) <= 5:
            melhor_bloco = candidatos
            break

        if len(candidatos) > 5:
            melhor_bloco = candidatos[:5]
            break

    alternativas = []
    for i in melhor_bloco:
        txt = bloco[i]["text"].strip()
        alternativas.append(
            {
                "texto": limpar_marcador_alternativa(txt),
                "correta": detectar_alternativa_correta(bloco[i]),
                "indice": i,
            }
        )

    return alternativas


def fallback_alternativas_documento_unico(bloco):
    candidatos = []
    coletando = False

    for i in range(len(bloco) - 1, -1, -1):
        txt = bloco[i]["text"].strip()

        if not txt:
            continue

        if eh_linha_alternativa_implicita(txt):
            candidatos.append(i)
            coletando = True
            continue

        if coletando:
            break

    candidatos = sorted(candidatos)

    alternativas = []
    for i in candidatos:
        txt = bloco[i]["text"].strip()
        alternativas.append(
            {
                "texto": limpar_marcador_alternativa(txt),
                "correta": detectar_alternativa_correta(bloco[i]),
                "indice": i,
            }
        )

    if len(alternativas) > 5:
        alternativas = alternativas[:5]

    return alternativas


def fallback_ultimas_linhas_unico(bloco):
    indices = []
    coletando = False

    for i in range(len(bloco) - 1, -1, -1):
        txt = bloco[i]["text"].strip()

        if not txt:
            continue

        if eh_justificativa(txt) or eh_linha_gabarito(txt) or eh_titulo_forum(txt) or eh_titulo_questoes(txt):
            continue

        if eh_cabecalho_unidade(txt) or eh_assunto(txt):
            continue

        if eh_alternativa_explica(txt) or eh_linha_alternativa_implicita(txt):
            indices.append(i)
            coletando = True
            if len(indices) == 5:
                break
            continue

        if coletando:
            break

    indices = sorted(indices)

    alternativas = []
    for i in indices:
        txt = bloco[i]["text"].strip()
        alternativas.append(
            {
                "texto": limpar_marcador_alternativa(txt),
                "correta": detectar_alternativa_correta(bloco[i]),
                "indice": i,
            }
        )

    if len(alternativas) in {4, 5}:
        return alternativas

    return []


def parse_questao_documento_unico(assunto, bloco, numero, unidade, letra_gabarito=None):
    comentario = ""
    letra_gabarito_inline = None

    for p in bloco:
        txt = p["text"].strip()
        if not letra_gabarito_inline and eh_linha_gabarito(txt):
            letra_gabarito_inline = extrair_letra_gabarito_inline(txt)
        elif eh_linha_comentario(txt):
            trecho = limpar_linha_comentario(txt)
            comentario = f"{comentario} {trecho}".strip() if comentario else trecho

    letra_gabarito_final = letra_gabarito or letra_gabarito_inline
    minimo_necessario = 2 if letra_gabarito_final else 4

    if letra_gabarito_final:
        alternativas = coletar_alternativas_explicitas_flexivel_unico(bloco, minimo=minimo_necessario)
    else:
        alternativas = coletar_alternativas_explicitas_unico(bloco)

    if len(alternativas) < minimo_necessario:
        alternativas = coletar_alternativas_por_janela_unico(bloco)

    if len(alternativas) < minimo_necessario:
        alternativas = fallback_alternativas_documento_unico(bloco)

    if len(alternativas) < minimo_necessario:
        alternativas = fallback_ultimas_linhas_unico(bloco)

    alt_indices = {a["indice"] for a in alternativas}

    enunciado_partes = []
    for i, p in enumerate(bloco):
        txt = p["text"].strip()

        if i in alt_indices:
            continue

        if eh_linha_gabarito(txt) or eh_linha_comentario(txt):
            continue

        if eh_subtitulo_tematica_numerado(txt):
            continue

        if i == 0 and eh_inicio_questao_numerada(txt) and not eh_subtitulo_tematica_numerado(txt):
            txt = limpar_numero_questao_inicial(txt)

        enunciado_partes.append(txt)

    enunciado_txt = "\n".join([x for x in enunciado_partes if x.strip()]).strip()

    if letra_gabarito_final and alternativas:
        idx_gabarito = ord(letra_gabarito_final.upper()) - ord("A")
        if 0 <= idx_gabarito < len(alternativas):
            for alt in alternativas:
                alt["correta"] = False
            alternativas[idx_gabarito]["correta"] = True

    corretas = [a for a in alternativas if a["correta"]]

    if not corretas and alternativas:
        alternativas[0]["correta"] = True
        corretas = [alternativas[0]]
        print(
            f"[AVISO] {unidade} | Questão {numero:02d}: "
            f"nenhuma alternativa marcada em vermelho/amarelo/negrito/gabarito. Assumindo primeira alternativa."
        )

    incorretas = [a for a in alternativas if not a["correta"]]
    titulo = inferir_titulo(enunciado_txt, numero, assunto=assunto)

    return {
        "titulo": titulo,
        "enunciado": enunciado_txt,
        "correta": corretas[0]["texto"] if corretas else "",
        "incorretas": [a["texto"] for a in incorretas],
        "justificativa": comentario,
        "modo": "arquivo_unico_bloco_gabarito_comentario",
        "qtd_alternativas": len(alternativas),
        "tipo": "Objetiva",
    }



# =========================
# MODO NOVO: FÓRUM + QUESTIONÁRIO + GABARITO
# =========================
def extrair_numero_forum(texto: str):
    """Extrai o número de títulos como 'Fórum 01' ou 'Forum I'."""
    txt = texto.strip()

    m = re.match(r"^f[oó]rum\s+0*(\d+)\s*$", txt, flags=re.IGNORECASE)
    if m:
        return int(m.group(1))

    m = re.match(r"^f[oó]rum\s+([IVXLCDM]+)\s*$", txt, flags=re.IGNORECASE)
    if m:
        return romano_para_int(m.group(1))

    return None


def extrair_numero_questionario(texto: str):
    """Extrai o número de títulos como 'Questionário 01' ou 'Questionario I'."""
    txt = texto.strip()

    m = re.match(r"^question[áa]rio\s+0*(\d+)\s*$", txt, flags=re.IGNORECASE)
    if m:
        return int(m.group(1))

    m = re.match(r"^question[áa]rio\s+([IVXLCDM]+)\s*$", txt, flags=re.IGNORECASE)
    if m:
        return romano_para_int(m.group(1))

    return None


def arquivo_tem_padrao_forum_questionario_gabarito(paragrafos) -> bool:
    """
    Detecta o novo modelo sem interferir nos modelos antigos.

    O padrão esperado é:
    - Fórum NN: questão discursiva da Unidade NN;
    - Questionário NN: questões objetivas da Unidade NN;
    - Gabarito/GABARITO: respostas do questionário imediatamente anterior.
    """
    tem_forum_numerado = any(extrair_numero_forum(p["text"]) for p in paragrafos)
    tem_questionario_numerado = any(extrair_numero_questionario(p["text"]) for p in paragrafos)
    tem_gabarito_textual = any(eh_titulo_gabarito(p["text"]) for p in paragrafos)

    return tem_forum_numerado and tem_questionario_numerado and tem_gabarito_textual


def extrair_gabarito_textual(linhas: list[str]) -> dict:
    """
    Extrai gabaritos em parágrafo comum, por exemplo:
    '1) A; 2) B; 3) C.'
    '1) D – art. 206 do CTN; 2) A – art. 203 do CTN;'

    Retorno:
    {
        1: {"letra": "A", "comentario": ""},
        2: {"letra": "B", "comentario": "art. ..."},
    }
    """
    texto = " ".join(l.strip() for l in linhas if l and l.strip())
    texto = re.sub(r"\s+", " ", texto).strip()

    if not texto:
        return {}

    # Normaliza travessões e hífens usados como separador de comentário.
    texto = texto.replace("—", "–")

    padrao = re.compile(
        r"(\d{1,3})\s*[\)\.]\s*([A-E])(?:\s*[–-]\s*([^;]+?))?(?=\s*;|\s*\.\s*$|\s*$)",
        flags=re.IGNORECASE,
    )

    resultado = {}
    for m in padrao.finditer(texto):
        numero = int(m.group(1))
        letra = m.group(2).upper()
        comentario = (m.group(3) or "").strip()
        comentario = re.sub(r"\s+", " ", comentario).strip(" .;–-")
        resultado[numero] = {"letra": letra, "comentario": comentario}

    # Fallback para textos sem separador final consistente.
    if not resultado:
        pares = re.findall(r"(\d{1,3})\s*[\)\.]\s*([A-E])\b", texto, flags=re.IGNORECASE)
        resultado = {
            int(numero): {"letra": letra.upper(), "comentario": ""}
            for numero, letra in pares
        }

    return resultado


def separar_por_forum_questionario_gabarito(paragrafos):
    """
    Agrupa o novo modelo em unidades lógicas.

    Importante: os fóruns podem aparecer todos no início do Word e os
    questionários depois. Por isso, o vínculo é feito pelo número:
    Fórum 01 -> Unidade 1; Questionário 01 -> Unidade 1.
    """
    dados = {}
    secao_atual = None
    unidade_atual = None
    buffer_gabarito = []

    def garantir_unidade(numero):
        if numero not in dados:
            dados[numero] = {
                "forum": [],
                "questoes": [],
                "gabarito": {},
            }
        return dados[numero]

    def finalizar_gabarito():
        nonlocal buffer_gabarito
        if secao_atual == "gabarito" and unidade_atual is not None and buffer_gabarito:
            linhas = [p["text"] for p in buffer_gabarito]
            garantir_unidade(unidade_atual)["gabarito"] = extrair_gabarito_textual(linhas)
            buffer_gabarito = []

    for p in paragrafos:
        txt = p["text"].strip()
        if not txt:
            continue

        numero_forum = extrair_numero_forum(txt)
        if numero_forum is not None:
            finalizar_gabarito()
            unidade_atual = numero_forum
            garantir_unidade(unidade_atual)
            secao_atual = "forum"
            continue

        numero_questionario = extrair_numero_questionario(txt)
        if numero_questionario is not None:
            finalizar_gabarito()
            unidade_atual = numero_questionario
            garantir_unidade(unidade_atual)
            secao_atual = "questoes"
            continue

        if eh_titulo_gabarito(txt):
            finalizar_gabarito()
            secao_atual = "gabarito"
            buffer_gabarito = []
            continue

        if unidade_atual is None:
            continue

        if secao_atual == "forum":
            garantir_unidade(unidade_atual)["forum"].append(p)
        elif secao_atual == "questoes":
            garantir_unidade(unidade_atual)["questoes"].append(p)
        elif secao_atual == "gabarito":
            buffer_gabarito.append(p)

    finalizar_gabarito()

    return [(f"Unidade {numero}", dados[numero]) for numero in sorted(dados)]


def eh_inicio_questao_word(paragrafo) -> bool:
    """Identifica questão por nível 0 da lista automática ou por numeração digitada."""
    txt = paragrafo["text"].strip()

    if paragrafo.get("ilvl") == 0:
        return True

    return bool(re.match(r"^\d{1,3}\s*[\)\.]\s+.+", txt))


def eh_alternativa_word(paragrafo) -> bool:
    """Identifica alternativa por nível 1 da lista automática ou por letra digitada."""
    txt = paragrafo["text"].strip()

    if paragrafo.get("ilvl") == 1:
        return True

    return bool(re.match(r"^[A-Ea-e]\s*[\)\.]\s+.+", txt))


def limpar_numero_questao_word(texto: str) -> str:
    return re.sub(r"^\d{1,3}\s*[\)\.]\s*", "", texto.strip()).strip()


def limpar_alternativa_word(texto: str) -> str:
    return re.sub(r"^[A-Ea-e]\s*[\)\.]\s*", "", texto.strip()).strip()


def separar_questoes_por_lista_word(paragrafos_questoes):
    """
    Separa questões objetivas do novo modelo.

    Usa preferencialmente a numeração automática do Word:
    - ilvl 0: questão;
    - ilvl 1: alternativa.

    Também mantém fallback para arquivos em que '1)' e 'a)' estejam digitados
    manualmente no texto.
    """
    questoes = []
    atual = None

    for p in paragrafos_questoes:
        txt = p["text"].strip()
        if not txt:
            continue

        if eh_inicio_questao_word(p):
            if atual:
                questoes.append(atual)

            atual = {
                "enunciado_partes": [limpar_numero_questao_word(txt)],
                "alternativas": [],
            }
            continue

        if eh_alternativa_word(p):
            if atual:
                atual["alternativas"].append(limpar_alternativa_word(txt))
            continue

        if atual:
            atual["enunciado_partes"].append(txt)

    if atual:
        questoes.append(atual)

    return questoes


def parse_questao_modelo_forum_questionario(bloco_questao: dict, numero: int, unidade: str, info_gabarito=None):
    enunciado_txt = "\n".join(
        parte.strip()
        for parte in bloco_questao.get("enunciado_partes", [])
        if parte and parte.strip()
    ).strip()

    alternativas_texto = [
        alt.strip()
        for alt in bloco_questao.get("alternativas", [])
        if alt and alt.strip()
    ]

    letra_gabarito = None
    comentario_gabarito = ""

    if isinstance(info_gabarito, dict):
        letra_gabarito = info_gabarito.get("letra")
        comentario_gabarito = info_gabarito.get("comentario", "") or ""
    elif isinstance(info_gabarito, str):
        letra_gabarito = info_gabarito

    alternativas = []
    for idx, texto_alt in enumerate(alternativas_texto):
        letra_atual = chr(ord("A") + idx)
        alternativas.append(
            {
                "texto": texto_alt,
                "correta": bool(letra_gabarito and letra_atual == letra_gabarito.upper()),
                "indice": idx,
            }
        )

    if letra_gabarito and alternativas:
        idx_gabarito = ord(letra_gabarito.upper()) - ord("A")
        if not (0 <= idx_gabarito < len(alternativas)):
            print(
                f"[AVISO] {unidade} | Objetiva {numero:02d}: "
                f"gabarito '{letra_gabarito}' fora do intervalo de {len(alternativas)} alternativas."
            )
    elif alternativas:
        # Segurança: no novo modelo o gabarito textual deve existir. Se não existir,
        # preserva o comportamento antigo de assumir a primeira alternativa, mas avisa.
        alternativas[0]["correta"] = True
        print(
            f"[AVISO] {unidade} | Objetiva {numero:02d}: "
            f"gabarito textual não encontrado. Assumindo primeira alternativa."
        )

    corretas = [a for a in alternativas if a["correta"]]
    incorretas = [a for a in alternativas if not a["correta"]]

    titulo = inferir_titulo(enunciado_txt, numero)

    return {
        "titulo": titulo,
        "enunciado": enunciado_txt,
        "correta": corretas[0]["texto"] if corretas else "",
        "incorretas": [a["texto"] for a in incorretas],
        "justificativa": comentario_gabarito,
        "modo": "forum_questionario_gabarito_word",
        "qtd_alternativas": len(alternativas),
        "tipo": "Objetiva",
    }


def processar_unidades_forum_questionario(unidades):
    saida_blocos = []

    for unidade, dados_unidade in unidades:
        contador_discursiva = 1
        contador_objetiva = 1

        forum = dados_unidade.get("forum", [])
        if forum:
            q_forum = parse_questao_aberta_forum(forum, contador_discursiva, unidade)
            if q_forum["enunciado"]:
                saida_blocos.append(montar_gift(unidade, q_forum))
                print(
                    f"[OK] {unidade} | Discursiva {contador_discursiva:02d} | "
                    f"modo=forum_numerado | tipo={q_forum['tipo']}"
                )
                contador_discursiva += 1

        questoes = separar_questoes_por_lista_word(dados_unidade.get("questoes", []))
        gabarito = dados_unidade.get("gabarito", {}) or {}

        for bloco_questao in questoes:
            info_gabarito = gabarito.get(contador_objetiva)
            q = parse_questao_modelo_forum_questionario(
                bloco_questao,
                contador_objetiva,
                unidade,
                info_gabarito=info_gabarito,
            )

            minimo_alternativas = 2 if info_gabarito else 4
            if q["enunciado"] and q["correta"] and q["qtd_alternativas"] >= minimo_alternativas:
                saida_blocos.append(montar_gift(unidade, q))
                print(
                    f"[OK] {unidade} | Objetiva {contador_objetiva:02d} | "
                    f"modo={q['modo']} | alternativas={q['qtd_alternativas']} | "
                    f"gabarito={info_gabarito.get('letra') if isinstance(info_gabarito, dict) else info_gabarito}"
                )
            else:
                print(
                    f"[AVISO] {unidade} | Objetiva {contador_objetiva:02d}: "
                    f"ignorada | alternativas={q['qtd_alternativas']} | modo={q['modo']}"
                )

            contador_objetiva += 1

        if gabarito and len(questoes) != len(gabarito):
            print(
                f"[AVISO] {unidade}: quantidade de questões extraídas ({len(questoes)}) "
                f"diferente da quantidade de itens no gabarito ({len(gabarito)})."
            )

    return saida_blocos


# =========================
# PROCESSAMENTO
# =========================
def _imagem_por_marcador(marcador: str):
    for imagem in IMAGENS_EXTRAIDAS.values():
        if imagem["marcador"] == marcador:
            return imagem
    return None


def _substituir_marcadores_html(texto: str, modo: str) -> tuple[str, list]:
    """Substitui marcadores por IMG externo (GIFT) ou @@PLUGINFILE@@ (XML)."""
    if not texto:
        return "", []

    usados = []
    resultado = texto
    for imagem in IMAGENS_EXTRAIDAS.values():
        marcador = imagem["marcador"]
        if marcador not in resultado:
            continue
        usados.append(imagem)
        nome_url = quote(imagem["nome"])
        if modo == "gift" and not URL_BASE_IMAGENS:
            # Na geração exclusiva de XML, preserva o marcador para que a
            # segunda etapa o converta em @@PLUGINFILE@@ com Base64.
            continue
        if modo == "gift":
            src = f"{URL_BASE_IMAGENS.rstrip('/')}/{nome_url}"
        else:
            src = f"@@PLUGINFILE@@/{nome_url}"
        tag = (
            f'<p class="imagem-questao"><img src="{src}" '
            f'alt="Imagem da questão" style="max-width:100%;height:auto;"></p>'
        )
        resultado = resultado.replace(marcador, tag)
    return resultado, usados


def _escapar_gift_preservando_imagens(texto: str) -> str:
    """Escapa texto GIFT, mas mantém as tags HTML geradas para imagens."""
    if not texto:
        return ""
    html_texto, imagens = _substituir_marcadores_html(texto, "gift")
    protegidos = {}
    for idx, match in enumerate(re.findall(r'<p class="imagem-questao">.*?</p>', html_texto, flags=re.DOTALL)):
        token = f"__HTML_PROTEGIDO_{idx}__"
        protegidos[token] = match
        html_texto = html_texto.replace(match, token, 1)
    html_texto = escapar_gift(html_texto)
    for token, trecho in protegidos.items():
        html_texto = html_texto.replace(token, trecho)
    return html_texto


def montar_gift(unidade: str, questao: dict) -> str:
    linhas = []
    linhas.append(f"// [tag: {unidade}]")
    linhas.append(f"// [tag: {questao.get('tipo', 'Objetiva')}]")
    linhas.append(f"// [tag: {TIPO_MATERIAL}]")
    linhas.append(f"// [tag: {DISCIPLINA}]")
    linhas.append(f"::{escapar_gift(questao['titulo'])}::")

    tem_imagem = any(
        img["marcador"] in (questao.get("enunciado", "") + questao.get("justificativa", ""))
        for img in IMAGENS_EXTRAIDAS.values()
    )
    prefixo = "[html]" if tem_imagem else ""
    linhas.append(prefixo + _escapar_gift_preservando_imagens(questao["enunciado"]))

    if questao.get("tipo") == "Discursiva":
        linhas.append("{}")
        return "\n".join(linhas)

    linhas.append("{")
    linhas.append(f"    ={escapar_gift(questao['correta'])}")

    for alt in questao["incorretas"]:
        linhas.append(f"    ~{escapar_gift(alt)}")

    if questao["justificativa"]:
        feedback = _escapar_gift_preservando_imagens(
            "Justificativa Geral: " + questao["justificativa"]
        )
        linhas.append(f"    #### {feedback}")

    linhas.append("}")
    return "\n".join(linhas)



# =========================
# MODO NOVO 2: QUESTIONÁRIOS + FÓRUM POR UNIDADE
# =========================
# Correções pontuais para questões sem marcação de resposta no .docx.
# Para reutilizar em outros arquivos, mantenha vazio ou ajuste conforme necessário.
CORRECOES_MANUAIS = {}

def eh_titulo_questionarios_geral(texto: str) -> bool:
    """Detecta cabeçalhos gerais como 'QUESTIONÁRIOS' ou 'QUESTIONÁRIO DAS UNIDADES'."""
    txt = normalizar_texto(texto)
    return txt in {"questionários", "questionarios", "questionário das unidades", "questionario das unidades"} or txt.startswith("questionário das unidades") or txt.startswith("questionario das unidades")


def extrair_numero_unidade_simples(texto: str):
    """
    Extrai cabeçalhos reais de unidade, inclusive quando há título depois do número:
    - UNIDADE 01 – TÍTULO
    - UNIDADE 01: TÍTULO
    - UNIDADE I – TÍTULO

    Não considera alternativas como "Unidade e educação infantil", pois exige
    número arábico ou romano logo depois da palavra UNIDADE.
    """
    txt = texto.strip()

    m = re.match(r"^Unidade\s+0*(\d+)\b(?:\s*[:\-–—].*)?$", txt, flags=re.IGNORECASE)
    if m:
        return int(m.group(1))

    m = re.match(r"^Unidade\s+([IVXLCDM]+)\b(?:\s*[:\-–—].*)?$", txt, flags=re.IGNORECASE)
    if m:
        return romano_para_int(m.group(1))

    return None


def eh_titulo_foruns_geral(texto: str) -> bool:
    """Detecta o cabeçalho geral de fóruns: FÓRUM, FÓRUNS, FORUM ou FORUNS."""
    txt = normalizar_texto(texto)
    return txt in {"fórum", "forum", "fóruns", "foruns"}


def limpar_prefixo_forum(texto: str) -> str:
    """Remove o prefixo visual 'FÓRUM:' quando ele vem colado ao enunciado."""
    return re.sub(r"^f[oó]rum\s*:\s*", "", texto.strip(), flags=re.IGNORECASE).strip()


def arquivo_tem_padrao_questionarios_forum_unidade(paragrafos) -> bool:
    """
    Detecta o padrão do arquivo novo enviado:
    - bloco 'QUESTIONÁRIOS';
    - unidades dentro desse bloco;
    - bloco 'FÓRUM';
    - unidades dentro do bloco de fórum.
    """
    tem_questionarios = any(eh_titulo_questionarios_geral(p["text"]) for p in paragrafos)
    tem_forum_geral = any(eh_titulo_foruns_geral(p["text"]) for p in paragrafos)
    qtd_unidades = sum(1 for p in paragrafos if extrair_numero_unidade_simples(p["text"]) is not None)
    return tem_questionarios and tem_forum_geral and qtd_unidades >= 2


def separar_questionarios_forum_por_unidade(paragrafos):
    """
    Separa o arquivo no formato:
    QUESTIONÁRIOS -> UNIDADE 1..N
    FÓRUM -> UNIDADE 1..N
    """
    dados = {}
    secao_atual = None
    unidade_atual = None

    def garantir_unidade(numero: int):
        if numero not in dados:
            dados[numero] = {"questoes": [], "forum": []}
        return dados[numero]

    for p in paragrafos:
        txt = p["text"].strip()
        if not txt:
            continue

        if eh_titulo_questionarios_geral(txt):
            secao_atual = "questoes"
            unidade_atual = None
            continue

        if eh_titulo_foruns_geral(txt):
            secao_atual = "forum"
            unidade_atual = None
            continue

        numero_unidade = extrair_numero_unidade_simples(txt)
        if numero_unidade is not None and secao_atual in {"questoes", "forum"}:
            unidade_atual = numero_unidade
            garantir_unidade(unidade_atual)
            continue

        if secao_atual in {"questoes", "forum"} and unidade_atual is not None:
            garantir_unidade(unidade_atual)[secao_atual].append(p)

    return [(f"Unidade {numero}", dados[numero]) for numero in sorted(dados)]


def detectar_num_id_principal_questoes(paragrafos_questoes):
    """
    No arquivo novo, o Word usa listas automáticas sem exibir '1.' ou 'a)'
    em p.text. A pergunta e as alternativas vêm com ilvl=0, mas com num_id
    diferente. O num_id principal é o que aparece uma vez para cada questão
    da unidade, normalmente 10 vezes.
    """
    from collections import Counter

    contagem = Counter(
        p.get("num_id")
        for p in paragrafos_questoes
        if p.get("num_id") is not None
    )

    if not contagem:
        return None

    return contagem.most_common(1)[0][0]


def separar_questoes_por_num_id_principal(paragrafos_questoes):
    num_id_principal = detectar_num_id_principal_questoes(paragrafos_questoes)
    if num_id_principal is None:
        return []

    questoes = []
    bloco = []

    for p in paragrafos_questoes:
        txt = p["text"].strip()
        if not txt:
            continue

        if p.get("num_id") == num_id_principal:
            if bloco:
                questoes.append(bloco)
            bloco = [p]
            continue

        if bloco:
            bloco.append(p)

    if bloco:
        questoes.append(bloco)

    return questoes


def localizar_grupo_final_de_alternativas(bloco_questao):
    """
    Localiza o último grupo de alternativas finais dentro da questão.

    O novo arquivo usa dois padrões misturados:
    1. listas automáticas do Word, em que as letras a), b), c) não aparecem
       em p.text, mas as alternativas compartilham o mesmo num_id;
    2. alternativas digitadas manualmente, como "a) texto", sem num_id.

    A função guarda todos os grupos válidos e retorna o último, evitando
    confundir enumerações internas I, II, III com as alternativas finais.
    """
    grupos = []
    i = 1

    while i < len(bloco_questao):
        txt = bloco_questao[i]["text"].strip()

        # Caso 1: alternativas digitadas manualmente: a), b), c), d), e).
        if eh_alternativa_explica(txt):
            grupo = [i]
            i += 1
            while i < len(bloco_questao) and eh_alternativa_explica(bloco_questao[i]["text"].strip()):
                grupo.append(i)
                i += 1
            if 2 <= len(grupo) <= 5:
                grupos.append(grupo)
            continue

        # Caso 2: alternativas em lista automática do Word.
        num_id = bloco_questao[i].get("num_id")
        if num_id is not None:
            grupo = [i]
            i += 1
            while i < len(bloco_questao) and bloco_questao[i].get("num_id") == num_id:
                grupo.append(i)
                i += 1
            if 2 <= len(grupo) <= 5:
                grupos.append(grupo)
            continue

        i += 1

    return grupos[-1] if grupos else []


def parse_questao_questionarios_por_unidade(bloco_questao, numero: int, unidade: str):
    alt_indices = set(localizar_grupo_final_de_alternativas(bloco_questao))

    alternativas = []
    for idx in sorted(alt_indices):
        p = bloco_questao[idx]
        texto_alt = limpar_marcador_alternativa(p["text"].strip())
        if not texto_alt:
            continue
        alternativas.append(
            {
                "texto": texto_alt,
                # Neste padrão, a resposta correta vem marcada por realce amarelo
                # ou cor vermelha. O negrito é usado em títulos/ênfases e não deve
                # ser interpretado como gabarito.
                "correta": any((run["red"] or run["yellow"]) and run["text"].strip() for run in p["runs"]),
                "indice": idx,
            }
        )

    enunciado_partes = []
    for idx, p in enumerate(bloco_questao):
        if idx in alt_indices:
            continue
        txt = p["text"].strip()
        if txt:
            enunciado_partes.append(txt)

    enunciado_txt = "\n".join(enunciado_partes).strip()

    letra_manual = CORRECOES_MANUAIS.get((unidade, numero))
    if letra_manual and alternativas:
        idx_manual = ord(letra_manual.upper()) - ord("A")
        if 0 <= idx_manual < len(alternativas):
            for alt in alternativas:
                alt["correta"] = False
            alternativas[idx_manual]["correta"] = True
            print(
                f"[INFO] {unidade} | Objetiva {numero:02d}: "
                f"correção manual aplicada: alternativa {letra_manual.upper()}."
            )

    corretas = [a for a in alternativas if a["correta"]]
    if len(corretas) > 1:
        print(
            f"[AVISO] {unidade} | Objetiva {numero:02d}: "
            f"mais de uma alternativa marcada como correta. Mantendo a primeira."
        )
        primeira = corretas[0]
        for alt in alternativas:
            alt["correta"] = alt is primeira
        corretas = [primeira]

    if not corretas and alternativas:
        alternativas[0]["correta"] = True
        corretas = [alternativas[0]]
        print(
            f"[AVISO] {unidade} | Objetiva {numero:02d}: "
            f"correta não encontrada por formatação. Assumindo primeira alternativa."
        )

    incorretas = [a for a in alternativas if not a["correta"]]

    return {
        "titulo": inferir_titulo(enunciado_txt, numero),
        "enunciado": enunciado_txt,
        "correta": corretas[0]["texto"] if corretas else "",
        "incorretas": [a["texto"] for a in incorretas],
        "justificativa": "",
        "modo": "questionarios_forum_por_unidade_num_id",
        "qtd_alternativas": len(alternativas),
        "tipo": "Objetiva",
    }


def parse_forum_questionarios_por_unidade(paragrafos_forum, numero: int, unidade: str):
    enunciado_txt = "\n".join(
        limpar_prefixo_forum(p["text"].strip())
        for p in paragrafos_forum
        if p["text"].strip()
    ).strip()

    titulo_base = inferir_titulo(enunciado_txt, numero)
    titulo = titulo_base.replace("Questão", "Fórum", 1)

    return {
        "titulo": titulo,
        "enunciado": enunciado_txt,
        "correta": "",
        "incorretas": [],
        "justificativa": "",
        "modo": "forum_por_unidade",
        "qtd_alternativas": 0,
        "tipo": "Discursiva",
    }


def processar_unidades_questionarios_forum_por_unidade(unidades):
    saida_blocos = []

    for unidade, dados_unidade in unidades:
        questoes = separar_questoes_por_num_id_principal(dados_unidade.get("questoes", []))
        contador_objetiva = 1

        for bloco_questao in questoes:
            q = parse_questao_questionarios_por_unidade(bloco_questao, contador_objetiva, unidade)

            if q["enunciado"] and q["correta"] and q["qtd_alternativas"] >= 2:
                saida_blocos.append(montar_gift(unidade, q))
                print(
                    f"[OK] {unidade} | Objetiva {contador_objetiva:02d} | "
                    f"modo={q['modo']} | alternativas={q['qtd_alternativas']}"
                )
            else:
                print(
                    f"[AVISO] {unidade} | Objetiva {contador_objetiva:02d}: "
                    f"ignorada | alternativas={q['qtd_alternativas']} | modo={q['modo']}"
                )

            contador_objetiva += 1

        forum = dados_unidade.get("forum", [])
        if forum:
            q_forum = parse_forum_questionarios_por_unidade(forum, 1, unidade)
            if q_forum["enunciado"]:
                saida_blocos.append(montar_gift(unidade, q_forum))
                print(
                    f"[OK] {unidade} | Discursiva 01 | "
                    f"modo={q_forum['modo']} | tipo={q_forum['tipo']}"
                )

        if len(questoes) not in {10, 20}:
            print(
                f"[AVISO] {unidade}: foram extraídas {len(questoes)} questões objetivas; "
                f"verifique se a unidade deveria ter 10 ou 20."
            )

    return saida_blocos

def arquivo_tem_padrao_template_questao_justificativa(paragrafos) -> bool:
    """
    Detecta o padrão Word tradicional dos arquivos enviados na auditoria:
    - cabeçalhos "Questão N";
    - alternativas digitadas como a), b), c), d), e);
    - linha de "Justificativa:" logo após as alternativas.

    Esse padrão precisa ser tratado antes do modo "estrutura nova", pois a
    presença de "Justificativa:" fazia o arquivo ser desviado indevidamente
    para um parser pensado para outro modelo de documento.
    """
    qtd_questoes = sum(1 for p in paragrafos if eh_cabecalho_questao(p["text"].strip()))
    qtd_justificativas = sum(1 for p in paragrafos if eh_linha_comentario(p["text"].strip()))
    qtd_alternativas = sum(1 for p in paragrafos if eh_alternativa_explica(p["text"].strip()))

    tem_identidade_template = any(
        "template para elaboração" in normalizar_texto(p["text"])
        or normalizar_texto(p["text"]) == "banco de questões"
        or normalizar_texto(p["text"]) == "banco de questoes"
        for p in paragrafos
    )

    return (
        qtd_questoes >= 1
        and qtd_justificativas >= 1
        and qtd_alternativas >= qtd_questoes * 4
        and (tem_identidade_template or qtd_questoes == qtd_justificativas)
    )


def processar_template_questao_justificativa(caminho_arquivo, unidade_padrao=None):
    """Processa arquivos Word no modelo Questão + alternativas + Justificativa."""
    paragrafos = extrair_paragrafos_com_runs(caminho_arquivo)
    unidade = extrair_unidade_do_nome_arquivo(caminho_arquivo) or unidade_padrao or "Unidade 1"
    questoes_brutas = separar_questoes(paragrafos)
    saida_blocos = []

    for contador, bloco in enumerate(questoes_brutas, start=1):
        q = parse_questao_modo_antigo(bloco, contador, unidade)

        if q["enunciado"] and q["correta"] and q["qtd_alternativas"] >= 4:
            q["modo"] = "template_word_questao_justificativa"
            saida_blocos.append(montar_gift(unidade, q))
            print(
                f"[OK] {unidade} | Questão {contador:02d} | "
                f"modo={q['modo']} | alternativas={q['qtd_alternativas']}"
            )
        else:
            print(
                f"[AVISO] {unidade} | Questão {contador:02d}: "
                f"ignorada | alternativas={q['qtd_alternativas']} | modo={q['modo']}"
            )

    if not questoes_brutas:
        print(f"[AVISO] {unidade}: nenhuma questão encontrada no padrão 'Questão N'.")

    return saida_blocos


def processar_unidades_estruturadas(caminho_arquivo, unidades, unidade_padrao=None):
    gabaritos = extrair_gabaritos_de_tabelas(caminho_arquivo)
    saida_blocos = []

    for idx_unidade, (unidade_detectada, bloco_unidade) in enumerate(unidades):
        unidade = unidade_detectada or unidade_padrao or f"Unidade {idx_unidade + 1}"
        forum, questoes_secao = dividir_secao_unidade(bloco_unidade)
        gabarito_unidade = gabaritos[idx_unidade] if idx_unidade < len(gabaritos) else {}

        contador_discursiva = 1
        contador_objetiva = 1

        blocos_forum = separar_forum_documento_unico(forum)
        for bloco_forum in blocos_forum:
            q = parse_questao_aberta_forum(bloco_forum, contador_discursiva, unidade)
            if q["enunciado"]:
                saida_blocos.append(montar_gift(unidade, q))
                print(
                    f"[OK] {unidade} | Discursiva {contador_discursiva:02d} | "
                    f"modo={q['modo']} | tipo={q['tipo']}"
                )
                contador_discursiva += 1

        questoes = separar_questoes_documento_unico(questoes_secao)
        for assunto, bloco_questao in questoes:
            letra_gabarito = gabarito_unidade.get(contador_objetiva)
            q = parse_questao_documento_unico(
                assunto,
                bloco_questao,
                contador_objetiva,
                unidade,
                letra_gabarito=letra_gabarito,
            )

            minimo_alternativas = 2 if letra_gabarito else 4
            if q["enunciado"] and q["correta"] and q["qtd_alternativas"] >= minimo_alternativas:
                saida_blocos.append(montar_gift(unidade, q))
                origem = "gabarito_tabela" if letra_gabarito else q["modo"]
                print(
                    f"[OK] {unidade} | Objetiva {contador_objetiva:02d} | "
                    f"modo={origem} | alternativas={q['qtd_alternativas']}"
                )
            else:
                print(
                    f"[AVISO] {unidade} | Objetiva {contador_objetiva:02d}: "
                    f"ignorada | alternativas={q['qtd_alternativas']} | modo={q['modo']}"
                )

            contador_objetiva += 1

    return saida_blocos


def arquivo_tem_estrutura_nova(paragrafos, caminho_arquivo) -> bool:
    if any(eh_cabecalho_unidade(p["text"].strip()) for p in paragrafos):
        return True

    if any(eh_titulo_forum(p["text"].strip()) for p in paragrafos):
        return True

    if any(eh_linha_gabarito(p["text"].strip()) for p in paragrafos):
        return True

    if any(eh_linha_comentario(p["text"].strip()) for p in paragrafos):
        return True

    if extrair_gabaritos_de_tabelas(caminho_arquivo):
        return True

    return False


def extrair_paragrafos_da_secao_forum(paragrafos):
    forum = []
    capturando = False

    for p in paragrafos:
        txt = p["text"].strip()

        if eh_titulo_forum(txt):
            capturando = True
            continue

        if capturando and txt:
            forum.append(p)

    return forum


def processar_arquivo_unico(caminho_arquivo):
    print(f"\n[INFO] Modo arquivo único: {caminho_arquivo}")
    paragrafos = extrair_paragrafos_com_runs(caminho_arquivo)

    if arquivo_tem_padrao_avaliacao_resposta_correta(paragrafos):
        print("[INFO] Padrão detectado: Avaliação + Resposta correta inline")
        return processar_avaliacao_resposta_correta(caminho_arquivo)

    if arquivo_tem_padrao_questionarios_forum_unidade(paragrafos):
        print("[INFO] Padrão detectado: QUESTIONÁRIOS + FÓRUM por UNIDADE")
        unidades = separar_questionarios_forum_por_unidade(paragrafos)
        return processar_unidades_questionarios_forum_por_unidade(unidades)

    if arquivo_tem_padrao_forum_questionario_gabarito(paragrafos):
        print("[INFO] Padrão detectado: Fórum + Questionário + Gabarito")
        unidades = separar_por_forum_questionario_gabarito(paragrafos)
        return processar_unidades_forum_questionario(unidades)

    if arquivo_tem_padrao_template_questao_justificativa(paragrafos):
        print("[INFO] Padrão detectado: Template Word - Questão + Alternativas + Justificativa")
        return processar_template_questao_justificativa(caminho_arquivo)

    unidades = separar_unidades_documento_unico(paragrafos)
    return processar_unidades_estruturadas(caminho_arquivo, unidades)


def processar_varios_arquivos(arquivos):
    print("\n[INFO] Modo vários arquivos.")
    saida_blocos = []

    for arquivo, unidade in arquivos:
        print(f"\nProcessando: {arquivo}")
        if not arquivo.exists():
            print(f"[ERRO] Arquivo não encontrado: {arquivo}")
            continue

        unidade_nome_arquivo = extrair_unidade_do_nome_arquivo(arquivo)
        unidade_final = unidade_nome_arquivo or unidade

        paragrafos = extrair_paragrafos_com_runs(arquivo)

        if arquivo_tem_padrao_avaliacao_resposta_correta(paragrafos):
            print("[INFO] Padrão detectado: Avaliação + Resposta correta inline")
            saida_blocos.extend(
                processar_avaliacao_resposta_correta(
                    arquivo,
                    unidade_padrao=unidade_final,
                )
            )
            continue

        if arquivo_tem_padrao_template_questao_justificativa(paragrafos):
            print("[INFO] Padrão detectado: Template Word - Questão + Alternativas + Justificativa")
            saida_blocos.extend(
                processar_template_questao_justificativa(
                    arquivo,
                    unidade_padrao=unidade_final,
                )
            )
            continue

        if arquivo_tem_estrutura_nova(paragrafos, arquivo):
            gabaritos = extrair_gabaritos_de_tabelas(arquivo)
            forum = extrair_paragrafos_da_secao_forum(paragrafos)
            if forum and not gabaritos:
                q = parse_questao_aberta_forum(forum, 1, unidade_final)
                if q["enunciado"]:
                    saida_blocos.append(montar_gift(unidade_final, q))
                    print(
                        f"[OK] {unidade_final} | Discursiva 01 | "
                        f"modo={q['modo']} | tipo={q['tipo']}"
                    )
                continue

            # Em modo de vários arquivos, cada .docx representa uma única unidade.
            # Cabeçalhos internos do tipo "Bloco X" devem funcionar apenas como
            # divisão temática, nunca como nova unidade.
            saida_blocos.extend(
                processar_unidades_estruturadas(
                    arquivo,
                    [(unidade_final, paragrafos)],
                    unidade_padrao=unidade_final,
                )
            )
            continue

        questoes_brutas = separar_questoes(paragrafos)

        contador = 1
        for bloco in questoes_brutas:
            q = parse_questao_modo_antigo(bloco, contador, unidade_final)

            if q["enunciado"] and q["correta"] and q["qtd_alternativas"] >= 4:
                saida_blocos.append(montar_gift(unidade_final, q))
                print(
                    f"[OK] {unidade_final} | Questão {contador:02d} | "
                    f"modo={q['modo']} | alternativas={q['qtd_alternativas']}"
                )
            else:
                print(
                    f"[AVISO] {unidade_final} | Questão {contador:02d}: "
                    f"ignorada | alternativas={q['qtd_alternativas']} | modo={q['modo']}"
                )

            contador += 1

    return saida_blocos


# =========================
# ENTRADAS DE MODO
# =========================
def escolher_modo_processamento():
    print("\nModo de processamento:")
    print("1 - Arquivo único com várias unidades")
    print("2 - Vários arquivos separados por unidade")

    opcao = input("Escolha uma opção (1 ou 2): ").strip()

    if opcao not in {"1", "2"}:
        raise ValueError("Escolha inválida. Digite 1 ou 2.")

    return opcao


def _mostrar_arquivos_word_disponiveis():
    disponiveis = listar_arquivos_word(BASE)
    if disponiveis:
        print("\nArquivos Word disponíveis na pasta (.docx e .doc):")
        for arquivo in disponiveis:
            print(f"  - {arquivo.name}")


def escolher_arquivo_unico():
    _mostrar_arquivos_word_disponiveis()
    while True:
        nome_arquivo = input("\nInforme o nome do arquivo único: ").strip()
        if not nome_arquivo:
            print("O nome do arquivo não pode ficar vazio. Tente novamente.")
            continue
        caminho = resolver_nome_arquivo_word(nome_arquivo)
        if caminho is None:
            print(f"[ERRO] Arquivo Word não encontrado: {nome_arquivo}. Use um arquivo .docx ou .doc.")
            continue
        try:
            return preparar_arquivo_word(caminho)
        except (RuntimeError, ValueError) as erro:
            print(f"[ERRO] {erro}")


def escolher_arquivos_por_unidade() -> list:
    _mostrar_arquivos_word_disponiveis()
    print("\nQuantas unidades deseja processar?")
    while True:
        try:
            qtd = int(input("Número de unidades: ").strip())
            if qtd < 1:
                raise ValueError
            break
        except ValueError:
            print("Digite um número válido maior que zero.")
    arquivos = []
    for i in range(1, qtd + 1):
        while True:
            nome = input(f"Nome do arquivo da Unidade {i}: ").strip()
            caminho = resolver_nome_arquivo_word(nome)
            if caminho is None:
                print(f"[ERRO] Arquivo Word não encontrado: {nome}. Use um arquivo .docx ou .doc.")
                continue
            try:
                arquivos.append((preparar_arquivo_word(caminho), f"Unidade {i}"))
                break
            except (RuntimeError, ValueError) as erro:
                print(f"[ERRO] {erro}")
    return arquivos



def escolher_preferencia_saida() -> str:
    """Solicita a preferência; a decisão final ocorre após analisar os Word."""
    print("\n" + "=" * 64)
    print("FORMATO DE SAÍDA")
    print("=" * 64)
    print("1 - Moodle XML")
    print("    Sempre gera XML, com ou sem imagens.")
    print("2 - GIFT")
    print("    Gera GIFT se não houver imagens; se houver, recomenda XML.")
    print("3 - Automático")
    print("    Sem imagens: GIFT | Com imagens: Moodle XML.")
    print("=" * 64)

    while True:
        opcao = input("Escolha uma opção (1, 2 ou 3): ").strip()
        if opcao == "1":
            return "xml"
        if opcao == "2":
            return "gift"
        if opcao == "3":
            return "automatico"
        print("[ERRO] Opção inválida. Digite 1, 2 ou 3.")


def _perguntar_sim_nao(pergunta: str) -> bool:
    while True:
        resposta = input(pergunta).strip().upper()
        if resposta in {"S", "SIM"}:
            return True
        if resposta in {"N", "NAO", "NÃO"}:
            return False
        print("[ERRO] Responda com S ou N.")


def resolver_formato_saida(preferencia: str, total_imagens: int) -> str:
    """Resolve o formato final depois de todos os Word serem analisados.

    Regras:
    - XML: sempre XML, com ou sem imagens.
    - GIFT sem imagens: GIFT direto.
    - GIFT com imagens: recomenda XML, pergunta pela troca e, se recusada,
      exige uma segunda confirmação antes de solicitar a URL pública.
    - Automático sem imagens: GIFT.
    - Automático com imagens: XML.
    """
    global FORMATO_SAIDA, URL_BASE_IMAGENS

    preferencia = (preferencia or "automatico").strip().lower()
    if preferencia not in {"xml", "gift", "automatico"}:
        raise ValueError(f"Preferência de saída inválida: {preferencia!r}")
    if total_imagens < 0:
        raise ValueError("O total de imagens não pode ser negativo.")

    # Evita reaproveitar URL ou formato de uma execução anterior/importação.
    FORMATO_SAIDA = ""
    URL_BASE_IMAGENS = ""

    # 1) XML escolhido: sempre XML.
    if preferencia == "xml":
        FORMATO_SAIDA = "xml"
        print("\n[INFO] Moodle XML selecionado pelo usuário.")
        if total_imagens > 0:
            print(
                f"[INFO] {total_imagens} imagem(ns) encontrada(s); "
                "todas serão incorporadas ao XML em Base64."
            )
        else:
            print(
                "[INFO] Nenhuma imagem encontrada, mas o arquivo será gerado "
                "em XML conforme a escolha do usuário."
            )
        return FORMATO_SAIDA

    # 2) Automático: decide sem novas perguntas.
    if preferencia == "automatico":
        if total_imagens > 0:
            FORMATO_SAIDA = "xml"
            print(
                f"\n[INFO] Modo automático: {total_imagens} imagem(ns) "
                "encontrada(s). Gerando Moodle XML."
            )
        else:
            FORMATO_SAIDA = "gift"
            print(
                "\n[INFO] Modo automático: nenhuma imagem encontrada. "
                "Gerando GIFT."
            )
        return FORMATO_SAIDA

    # 3) GIFT escolhido e nenhum arquivo contém imagem.
    if total_imagens == 0:
        FORMATO_SAIDA = "gift"
        print(
            "\n[INFO] GIFT selecionado e nenhuma imagem foi encontrada. "
            "Gerando GIFT diretamente."
        )
        return FORMATO_SAIDA

    # 4) GIFT escolhido, mas há imagens: recomenda troca para XML.
    print("\n" + "!" * 72)
    print(f"[ATENÇÃO] Foram encontradas {total_imagens} imagem(ns) nos documentos.")
    print("[ATENÇÃO] O GIFT não incorpora os arquivos de imagem no próprio .txt.")
    print("[RECOMENDAÇÃO] Troque para Moodle XML para importar tudo em um arquivo.")
    print("!" * 72)

    trocar_para_xml = _perguntar_sim_nao(
        "Deseja trocar o formato de GIFT para Moodle XML? [S/N]: "
    )
    if trocar_para_xml:
        FORMATO_SAIDA = "xml"
        print("[INFO] Formato alterado para Moodle XML.")
        return FORMATO_SAIDA

    # Segunda confirmação de segurança.
    print(
        "\n[ATENÇÃO] Ao permanecer no GIFT, as imagens deverão ser hospedadas "
        "externamente em uma URL pública e permanente."
    )
    continuar_gift = _perguntar_sim_nao(
        "Tem certeza de que deseja continuar em GIFT com imagens externas? [S/N]: "
    )
    if not continuar_gift:
        FORMATO_SAIDA = "xml"
        print("[INFO] Operação protegida: formato alterado para Moodle XML.")
        return FORMATO_SAIDA

    # Só solicita URL quando o usuário confirmou duas vezes que quer GIFT.
    while True:
        url = input(
            "Informe a URL pública da pasta de imagens "
            "(ex.: https://ava.exemplo.br/imagens_questoes): "
        ).strip().rstrip("/")

        if re.fullmatch(r"https?://[^\s]+", url, flags=re.IGNORECASE):
            URL_BASE_IMAGENS = url
            break

        print(
            "[ERRO] URL inválida. Informe um endereço completo iniciado por "
            "http:// ou https://, sem espaços."
        )

    FORMATO_SAIDA = "gift"
    print(
        "[INFO] GIFT será gerado com referências externas para as imagens em: "
        f"{URL_BASE_IMAGENS}"
    )
    return FORMATO_SAIDA


def preparar_blocos_gift_para_saida(saida_blocos):
    """Insere as tags IMG nos blocos que ainda possuem marcadores internos."""
    if not IMAGENS_EXTRAIDAS or not URL_BASE_IMAGENS:
        return saida_blocos

    preparados = []
    for bloco in saida_blocos:
        novo = bloco
        for imagem in IMAGENS_EXTRAIDAS.values():
            marcador = imagem["marcador"]
            if marcador not in novo:
                continue
            src = f"{URL_BASE_IMAGENS.rstrip('/')}/{quote(imagem['nome'])}"
            tag = (
                f'<p class="imagem-questao"><img src="{src}" '
                f'alt="Imagem da questão" style="max-width:100%;height:auto;"></p>'
            )
            novo = novo.replace(marcador, tag)
        preparados.append(novo)
    return preparados

def preparar_saidas():
    for caminho in (SAIDA_GIFT, SAIDA_XML):
        if caminho.exists():
            try:
                caminho.unlink()
                print(f"[INFO] Arquivo anterior removido: {caminho}")
            except Exception as e:
                raise PermissionError(f"Não foi possível remover {caminho}: {e}")

    if PASTA_IMAGENS.exists():
        for arquivo in PASTA_IMAGENS.iterdir():
            if arquivo.is_file():
                arquivo.unlink()
    else:
        PASTA_IMAGENS.mkdir(parents=True, exist_ok=True)


def salvar_imagens_extraidas():
    if not IMAGENS_EXTRAIDAS:
        print("[INFO] Nenhuma imagem incorporada foi encontrada nos documentos.")
        return

    PASTA_IMAGENS.mkdir(parents=True, exist_ok=True)
    for imagem in IMAGENS_EXTRAIDAS.values():
        destino = PASTA_IMAGENS / imagem["nome"]
        destino.write_bytes(imagem["dados"])
    print(f"[OK] {len(IMAGENS_EXTRAIDAS)} imagem(ns) exportada(s) para: {PASTA_IMAGENS}")


def _desescapar_gift(texto: str) -> str:
    """Desfaz apenas os escapes produzidos por escapar_gift()."""
    resultado = []
    i = 0
    while i < len(texto):
        if texto[i] == "\\" and i + 1 < len(texto):
            resultado.append(texto[i + 1])
            i += 2
        else:
            resultado.append(texto[i])
            i += 1
    return "".join(resultado)


def _parsear_bloco_gift_gerado(bloco: str) -> dict:
    linhas = bloco.splitlines()
    tags = []
    while linhas and linhas[0].startswith("// [tag:"):
        tag = linhas.pop(0).split(":", 1)[1].rsplit("]", 1)[0].strip()
        tags.append(tag)

    if not linhas or not linhas[0].startswith("::"):
        raise ValueError("Bloco GIFT sem título reconhecível.")

    titulo_linha = linhas.pop(0)
    titulo = _desescapar_gift(titulo_linha[2:titulo_linha.rfind("::")])

    try:
        idx_abre = linhas.index("{")
    except ValueError:
        idx_abre = len(linhas)

    enunciado = "\n".join(linhas[:idx_abre]).strip()
    if enunciado.startswith("[html]"):
        enunciado = enunciado[len("[html]"):]
    enunciado = _desescapar_gift(enunciado)

    tipo = "Discursiva" if idx_abre < len(linhas) and linhas[idx_abre:idx_abre + 2] == ["{", "}"] else "Objetiva"
    correta = ""
    incorretas = []
    justificativa = ""

    if idx_abre < len(linhas):
        for linha in linhas[idx_abre + 1:]:
            limpa = linha.strip()
            if limpa == "}":
                break
            if limpa.startswith("####"):
                justificativa = _desescapar_gift(limpa[4:].strip())
            elif limpa.startswith("="):
                correta = _desescapar_gift(limpa[1:].strip())
            elif limpa.startswith("~"):
                incorretas.append(_desescapar_gift(limpa[1:].strip()))

    return {
        "titulo": titulo,
        "enunciado": enunciado,
        "tipo": tipo,
        "correta": correta,
        "incorretas": incorretas,
        "justificativa": justificativa,
        "tags": tags,
    }


def _adicionar_texto(parent, nome, valor):
    elemento = ET.SubElement(parent, nome)
    texto = ET.SubElement(elemento, "text")
    texto.text = valor or ""
    return elemento


def _adicionar_arquivos_referenciados(parent, texto_html: str):
    for imagem in IMAGENS_EXTRAIDAS.values():
        if imagem["marcador"] not in texto_html:
            continue
        arquivo = ET.SubElement(
            parent,
            "file",
            {
                "name": imagem["nome"],
                "path": "/",
                "encoding": "base64",
            },
        )
        arquivo.text = imagem["base64"]


def gerar_moodle_xml(saida_blocos):
    quiz = ET.Element("quiz")

    # Categoria inicial opcional, facilita a organização após importação.
    categoria = ET.SubElement(quiz, "question", {"type": "category"})
    cat = ET.SubElement(categoria, "category")
    ET.SubElement(cat, "text").text = f"$course$/top/{DISCIPLINA}"

    for bloco in saida_blocos:
        dados = _parsear_bloco_gift_gerado(bloco)

        if dados["tipo"] == "Discursiva":
            q = ET.SubElement(quiz, "question", {"type": "essay"})
        else:
            q = ET.SubElement(quiz, "question", {"type": "multichoice"})

        _adicionar_texto(q, "name", dados["titulo"])

        # Troca eventual URL externa do GIFT pelo marcador interno antes de gerar XML.
        enunciado_original = dados["enunciado"]
        for imagem in IMAGENS_EXTRAIDAS.values():
            url_externa = f"{URL_BASE_IMAGENS.rstrip('/')}/{quote(imagem['nome'])}" if URL_BASE_IMAGENS else ""
            if url_externa:
                enunciado_original = enunciado_original.replace(url_externa, imagem["marcador"])
        enunciado_html, _ = _substituir_marcadores_html(enunciado_original, "xml")

        questiontext = ET.SubElement(q, "questiontext", {"format": "html"})
        ET.SubElement(questiontext, "text").text = enunciado_html.replace("\n", "<br>")
        _adicionar_arquivos_referenciados(questiontext, enunciado_original)

        justificativa_original = dados["justificativa"]
        for imagem in IMAGENS_EXTRAIDAS.values():
            url_externa = f"{URL_BASE_IMAGENS.rstrip('/')}/{quote(imagem['nome'])}" if URL_BASE_IMAGENS else ""
            if url_externa:
                justificativa_original = justificativa_original.replace(url_externa, imagem["marcador"])
        justificativa_html, _ = _substituir_marcadores_html(justificativa_original, "xml")
        generalfeedback = ET.SubElement(q, "generalfeedback", {"format": "html"})
        ET.SubElement(generalfeedback, "text").text = justificativa_html.replace("\n", "<br>")
        _adicionar_arquivos_referenciados(generalfeedback, justificativa_original)
        ET.SubElement(q, "defaultgrade").text = "1.0000000"
        ET.SubElement(q, "penalty").text = "0.3333333"
        ET.SubElement(q, "hidden").text = "0"
        ET.SubElement(q, "idnumber").text = ""

        if dados["tipo"] == "Discursiva":
            ET.SubElement(q, "responseformat").text = "editor"
            ET.SubElement(q, "responserequired").text = "1"
            ET.SubElement(q, "responsefieldlines").text = "15"
            ET.SubElement(q, "attachments").text = "0"
            ET.SubElement(q, "attachmentsrequired").text = "0"
        else:
            ET.SubElement(q, "single").text = "true"
            ET.SubElement(q, "shuffleanswers").text = "true"
            ET.SubElement(q, "answernumbering").text = "abc"
            ET.SubElement(q, "showstandardinstruction").text = "0"

            alternativas = [(dados["correta"], "100")] + [
                (alt, "0") for alt in dados["incorretas"]
            ]
            for alternativa, fracao in alternativas:
                answer = ET.SubElement(q, "answer", {"fraction": fracao, "format": "html"})
                ET.SubElement(answer, "text").text = alternativa
                _adicionar_texto(answer, "feedback", "")

        if dados["tags"]:
            tags_el = ET.SubElement(q, "tags")
            for tag in dados["tags"]:
                tag_el = ET.SubElement(tags_el, "tag")
                ET.SubElement(tag_el, "text").text = tag

    arvore = ET.ElementTree(quiz)
    ET.indent(arvore, space="  ")
    arvore.write(SAIDA_XML, encoding="utf-8", xml_declaration=True)
    print(f"[OK] Moodle XML gerado em: {SAIDA_XML}")


def main():
    global PREFERENCIA_SAIDA

    print("BASE:", BASE)

    # A preferência é coletada agora, mas somente será resolvida depois que
    # todos os documentos forem lidos e as imagens tiverem sido contabilizadas.
    PREFERENCIA_SAIDA = escolher_preferencia_saida()
    preparar_saidas()
    modo = escolher_modo_processamento()

    if modo == "1":
        caminho_arquivo = escolher_arquivo_unico()
        saida_blocos = processar_arquivo_unico(caminho_arquivo)
    else:
        arquivos = escolher_arquivos_por_unidade()
        saida_blocos = processar_varios_arquivos(arquivos)

    if not saida_blocos:
        print("[AVISO] Nenhuma questão válida foi extraída.")
        return

    total_imagens = len(IMAGENS_EXTRAIDAS)
    print(f"\n[DIAGNÓSTICO] Questões extraídas: {len(saida_blocos)}")
    print(f"[DIAGNÓSTICO] Imagens encontradas: {total_imagens}")

    formato_final = resolver_formato_saida(PREFERENCIA_SAIDA, total_imagens)

    if formato_final == "gift":
        blocos_gift = preparar_blocos_gift_para_saida(saida_blocos)
        SAIDA_GIFT.write_text("\n\n".join(blocos_gift), encoding="utf-8")
        print(f"[OK] GIFT gerado em: {SAIDA_GIFT}")

        if total_imagens:
            salvar_imagens_extraidas()
            print(
                "[ATENÇÃO] Publique os arquivos da pasta imagens_questoes na URL: "
                f"{URL_BASE_IMAGENS}"
            )
            print(
                "[ATENÇÃO] Não altere os nomes das imagens, pois eles já estão "
                "referenciados no arquivo GIFT."
            )
    else:
        # No XML, as imagens são incorporadas ao próprio arquivo em Base64.
        gerar_moodle_xml(saida_blocos)
        if total_imagens:
            print("[INFO] As imagens foram incorporadas ao XML; não é necessário hospedá-las.")


if __name__ == "__main__":
    main()
